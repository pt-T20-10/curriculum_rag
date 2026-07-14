from __future__ import annotations

from pathlib import Path
import asyncio
import re
import shutil
import subprocess
from types import SimpleNamespace
from zipfile import ZipFile

import pytest

from app.services.textbook import publisher
from app.services.textbook.reviewer import ReviewerAgent
from app.services.textbook import orchestrator, workflow_runner


def _state(title: str = "Lập trình C# Cơ bản") -> dict:
    return {
        "request": title,
        "textbook_title": title,
        "final_content": "# CHƯƠNG 1: MỞ ĐẦU\n\nNội dung kiểm thử.",
        "current_content": "",
        "preface_content": "",
        "enable_images": False,
        "language": "vi",
        "export_formats": ["PDF", "Word"],
    }


def _fake_pandoc(*, fail_pdf: bool, fail_docx: bool):
    def convert_file(source, *, to, outputfile, extra_args):
        if to == "pdf" and fail_pdf:
            raise RuntimeError("synthetic PDF failure")
        if to == "docx" and fail_docx:
            raise RuntimeError("synthetic DOCX failure")
        payload = b"%PDF-1.7\n" if to == "pdf" else b"PK\x03\x04DOCX"
        Path(outputfile).write_bytes(payload)
        return ""

    return SimpleNamespace(convert_file=convert_file)


@pytest.mark.parametrize(
    ("fail_pdf", "fail_docx", "has_pdf", "has_docx"),
    [
        (False, False, True, True),
        (True, False, False, True),
        (False, True, True, False),
        (True, True, False, False),
    ],
)
def test_publisher_keeps_export_artifacts_distinct(
    tmp_path: Path,
    monkeypatch,
    fail_pdf: bool,
    fail_docx: bool,
    has_pdf: bool,
    has_docx: bool,
) -> None:
    monkeypatch.setattr(publisher, "BASE_DIR", tmp_path)
    monkeypatch.setattr(publisher, "image_dir", tmp_path / "outputs" / "images")
    monkeypatch.setattr(
        publisher,
        "pypandoc",
        _fake_pandoc(fail_pdf=fail_pdf, fail_docx=fail_docx),
    )
    monkeypatch.setattr(publisher, "_get_word_reference_doc", lambda path: None)

    result = publisher.publish_curriculum(_state()) #type: ignore

    assert bool(result["final_pdf_filepath"]) is has_pdf
    assert bool(result["final_docx_filepath"]) is has_docx
    assert result["final_filepath"] == result["final_pdf_filepath"]
    assert result["final_markdown_filepath"].endswith(".md")
    if result["final_pdf_filepath"]:
        assert result["final_pdf_filepath"].endswith(".pdf")
    if result["final_docx_filepath"]:
        assert result["final_docx_filepath"].endswith(".docx")
    assert ("pdf" in result["export_errors"]) is fail_pdf
    assert ("docx" in result["export_errors"]) is fail_docx


def test_publisher_filename_sanitizes_control_whitespace(
    tmp_path: Path,
    monkeypatch,
) -> None:
    title = "LẬP TRÌNH WEB Hệ đào tạo Đại học\t_\nNgành đào tạo C++"
    monkeypatch.setattr(publisher, "BASE_DIR", tmp_path)
    monkeypatch.setattr(publisher, "image_dir", tmp_path / "outputs" / "images")
    monkeypatch.setattr(
        publisher,
        "pypandoc",
        _fake_pandoc(fail_pdf=False, fail_docx=False),
    )
    monkeypatch.setattr(publisher, "_get_word_reference_doc", lambda path: None)

    result = publisher.publish_curriculum(_state(title)) #type: ignore

    markdown_path = Path(result["final_markdown_filepath"])
    assert markdown_path.is_file()
    assert "\t" not in markdown_path.name
    assert "\n" not in markdown_path.name
    assert result["final_pdf_filepath"]
    assert result["final_docx_filepath"]


def test_preview_cli_exports_pdf_and_docx_to_output_dir(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from scripts import convert_markdown_to_pdf_docx

    md_path = tmp_path / "sample.md"
    md_path.write_text(
        "---\n"
        'title-meta: "Giáo trình kiểm thử"\n'
        "---\n\n"
        "# Lời nói đầu\n\n"
        "Nội dung mở đầu.\n\n"
        "# CHƯƠNG 1: MỞ ĐẦU\n\n"
        "Nội dung chương.",
        encoding="utf-8",
    )
    output_dir = tmp_path / "preview"

    monkeypatch.setattr(
        convert_markdown_to_pdf_docx.publisher,
        "pypandoc",
        _fake_pandoc(fail_pdf=False, fail_docx=False),
    )
    monkeypatch.setattr(
        convert_markdown_to_pdf_docx.publisher,
        "_get_word_reference_doc",
        lambda path: None,
    )
    monkeypatch.setattr(
        convert_markdown_to_pdf_docx.publisher,
        "finalize_word_docx",
        lambda *args, **kwargs: True,
    )

    artifacts = convert_markdown_to_pdf_docx.convert_markdown_to_pdf_docx(
        md_path,
        output_dir,
        "vi",
    )

    assert artifacts["pdf"] == output_dir / "sample_preview.pdf"
    assert artifacts["word"] == output_dir / "sample_preview.docx"
    assert artifacts["pdf"].read_bytes().startswith(b"%PDF")
    assert artifacts["word"].read_bytes().startswith(b"PK")


def test_add_figure_numbers_uses_chapter_scoped_sequence() -> None:
    content = (
        "![Ảnh bìa](cover.png){width=70%}\n\n"
        "# CHƯƠNG 1: MỞ ĐẦU\n\n"
        "## 1.1 Khái niệm\n\n"
        "![Sơ đồ thứ nhất](one.png){width=70%}\n\n"
        "### 1.1.1 Chi tiết\n\n"
        "![Hình 1.1.2.1: Sơ đồ thứ hai](two.png)\n\n"
        "## 1.2 Ứng dụng\n\n"
        "![Sơ đồ thứ ba](three.png)\n\n"
        "# CHƯƠNG 2: NÂNG CAO\n\n"
        "![Sơ đồ chương hai](four.png)\n"
    )

    numbered = publisher.add_figure_numbers(content, language="vi")

    assert "![Ảnh bìa](cover.png){width=70%}" in numbered
    assert "![Hình 1.1: Sơ đồ thứ nhất](one.png){width=70%}" in numbered
    assert "![Hình 1.2: Sơ đồ thứ hai](two.png)" in numbered
    assert "![Hình 1.3: Sơ đồ thứ ba](three.png)" in numbered
    assert "![Hình 2.1: Sơ đồ chương hai](four.png)" in numbered
    assert publisher.add_figure_numbers(numbered, language="vi") == numbered


def test_add_figure_numbers_supports_english_and_word_figure_list() -> None:
    content = (
        "# Preface\n\n"
        "Introductory text.\n\n"
        "# CHAPTER 3: BASICS\n\n"
        "## 3.1 Concepts\n\n"
        "![First diagram](one.png)\n\n"
        "### 3.1.1 Details\n\n"
        "![Figure 3.1.1.4: Second diagram](two.png)\n"
    )

    numbered = publisher.add_figure_numbers(content, language="en")
    word_md = publisher._prepare_word_md(
        numbered,
        title="Sample Book",
        enable_images=True,
        language="en",
    )

    assert "![Figure 3.1: First diagram](one.png)" in numbered
    assert "![Figure 3.2: Second diagram](two.png)" in numbered
    assert "# List of Figures" not in word_md
    assert re.search(r"(?m)^List of Figures$", word_md)
    assert publisher._WORD_FIGURE_LIST_MARKER in word_md
    assert "- Figure 3.1: First diagram" not in word_md


def test_word_markdown_spacing_normalizes_lists_and_glued_math() -> None:
    content = (
        "Đoạn trước có mảng$M \\times N$chứa giá trị.\n"
        "Giá trị $ C_{total} $ dùng với $ M $ và $N$liền chữ.\n"
        "Ví dụ $L = 0,1$ H,$C = 100$ F.\n"
        "- Công suất điện: $P = U \\cdot I$.\n"
        "- Năng lượng điện: $W = P \\cdot t$ với $t$ là thời gian.\n"
        "- Công suất điều hòa: `P_{đh} = 2,2` kW và P$_0$ = 15 W/m$^2$.\n"
        "$$\nS = \\frac{P_{đh}}{cosφ}\n$$\n"
        "Các đại lượng gồm số hàng ($M$) và số cột ($N$).\n"
        "Sobel theo chiều ngang ($G_x$) và dọc ($G_y$).\n"
        "Lỗi cũ có ($G_x $) và dọc ($ G_y$).\n"
        "- Công suất biểu kiến: $S = U \\cdot I$ Trong đó $\\varphi$ là góc lệch pha.\n"
        "Trong đó:\n"
        "- $M$: chiều rộng\n"
        "- $N$: chiều cao\n"
        "Đoạn sau."
    )

    normalized = publisher.fix_math_formatting(content)
    normalized = publisher.normalize_math_identifier_formatting(normalized)
    normalized = publisher.normalize_formula_explanations(normalized)
    normalized = publisher.normalize_markdown_list_spacing(normalized)

    assert "mảng $M \\times N$ chứa" in normalized
    assert "$C_{total}$" in normalized
    assert "$M$" in normalized
    assert "$N$ liền chữ" in normalized
    assert "$L = 0,1$ H, $C = 100$ F" in normalized
    assert "$P = U \\cdot I$.\n- Năng lượng" in normalized
    assert "$P = U \\cdot I$. - Năng lượng" not in normalized
    assert "$P_{dh} = 2,2$" in normalized
    assert "$P_{đh}" not in normalized
    assert "`P_{đh}" not in normalized
    assert "$P_{0}$ = 15" in normalized
    assert "W/m$^2$" in normalized
    assert "\\frac{P_{dh}}{\\cos\\varphi}" in normalized
    assert "$ C_{total} $" not in normalized
    assert "số hàng ($M$) và số cột ($N$)" in normalized
    assert "ngang ($G_x$) và dọc ($G_y$)" in normalized
    assert "$G_x $" not in normalized
    assert "$ G_y$" not in normalized
    assert "vàsố" not in normalized
    assert "- Công suất biểu kiến: $S = U \\cdot I$\n\nTrong đó:\n\n- $\\varphi$: góc lệch pha." in normalized
    assert "Trong đó:\n\n- $M$: chiều rộng" in normalized
    assert "- $N$: chiều cao\n\nĐoạn sau." in normalized


def test_markdown_image_blocks_are_standalone_for_word_captions() -> None:
    content = (
        "Đoạn trước ![Sơ đồ mạch](outputs/images/a.png){width=70%} Đoạn sau.\n\n"
        "> \n\n"
        "![Ảnh thứ hai](outputs/images/b.png){width=70%}\n"
    )

    normalized = publisher.normalize_markdown_image_blocks(content)
    numbered = publisher.add_figure_numbers(
        "# CHƯƠNG 1: MỞ ĐẦU\n\n" + normalized,
        language="vi",
    )

    assert "Đoạn trước\n\n![Sơ đồ mạch](outputs/images/a.png){width=70%}\n\nĐoạn sau." in normalized
    assert ">" not in normalized
    assert "![Hình 1.1: Sơ đồ mạch](outputs/images/a.png){width=70%}" in numbered
    assert "![Hình 1.2: Ảnh thứ hai](outputs/images/b.png){width=70%}" in numbered


def test_partial_markdown_checkpoint_preserves_reviewed_content(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(publisher, "BASE_DIR", tmp_path)
    state = {
        **_state("Giáo trình Web"),
        "rag_collection_name": "dynamic_context_221_20260704",
        "final_content": "# CHƯƠNG 1: HTML\n\nNội dung đã hoàn thành.",
        "current_content": "## 1.2 CSS\n\nNội dung vừa qua reviewer và illustrator.",
    }

    path = publisher.save_partial_markdown_checkpoint(state)

    assert path is not None
    checkpoint = Path(path)
    assert checkpoint.is_file()
    content = checkpoint.read_text(encoding="utf-8")
    assert "Auto-saved checkpoint" in content
    assert "Nội dung đã hoàn thành." in content
    assert "Nội dung vừa qua reviewer và illustrator." in content
    assert checkpoint.name.endswith(".partial.md")


def test_context_route_runs_targeted_recovery_once(monkeypatch) -> None:
    monkeypatch.setattr(
        orchestrator.settings,
        "CRAG_MAX_CONTEXT_RETRIES",
        2,
        raising=False,
    )
    monkeypatch.setattr(
        orchestrator.settings,
        "CRAG_TARGETED_RECOVERY_ENABLED",
        True,
        raising=False,
    )

    state = {
        "context_quality": "insufficient",
        "rag_retrieval_attempts": 2,
        "used_rag_queries": ["q1", "q2"],
        "rag_recovery_attempted": False,
    }

    assert (
        orchestrator.route_after_context_evaluation(state) #type: ignore
        == orchestrator.WorkflowDecision.RECOVER_CONTEXT
    )

    state["rag_recovery_attempted"] = True
    assert (
        orchestrator.route_after_context_evaluation(state) #type: ignore
        == orchestrator.WorkflowDecision.FINISHED
    )


@pytest.mark.parametrize(
    "title",
    [
        "Lập trình C# Cơ bản",
        'Tiêu đề có "dấu nháy"',
        "Mảng [A] và đường dẫn C:\\Temp",
        "Giáo trình tiếng Việt",
    ],
)
def test_typst_title_block_uses_a_string_literal(title: str) -> None:
    block = publisher._build_title_block(title)
    assert "#text(" in block
    assert f")[{title}]" not in block
    assert publisher._typst_string_literal(title) in block


def test_finalize_word_docx_localizes_toc_fonts_and_footer(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE

    path = tmp_path / "book.docx"
    doc = docx.Document()
    if "Source Code" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    title = doc.add_paragraph("Giáo trình kiểm thử")
    title.style = doc.styles["Title"]
    toc = doc.add_paragraph("Table of Contents")
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(qn("w:val"), "TOCHeading")
    toc._p.get_or_add_pPr().append(pstyle)
    doc.add_heading("Danh mục hình", level=1)
    doc.add_paragraph(publisher._WORD_FIGURE_LIST_MARKER)
    doc.add_heading("Lời nói đầu", level=1)
    doc.add_paragraph("Nội dung lời nói đầu cần được căn đều như bản PDF.")
    doc.add_heading("CHƯƠNG 1: MỞ ĐẦU", level=1)
    doc.add_heading("1.1 Khái niệm", level=2)
    doc.add_heading("1.1.1 Chi tiết", level=3)
    code = doc.add_paragraph("print('left aligned')")
    code.style = doc.styles["Source Code"]
    drawing_paragraph = doc.add_paragraph()
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:drawing"))
    drawing_paragraph._p.append(run)
    caption = doc.add_paragraph("Hình 1.1: Sơ đồ kiểm thử")
    caption.style = doc.styles["Caption"]
    doc.add_paragraph("Hình 1.2: Caption không có style riêng")
    missing_caption_drawing = doc.add_paragraph()
    missing_run = OxmlElement("w:r")
    missing_drawing = OxmlElement("w:drawing")
    missing_inline = OxmlElement("wp:inline")
    doc_pr = OxmlElement("wp:docPr")
    doc_pr.set("id", "9")
    doc_pr.set("name", "Picture 9")
    doc_pr.set("descr", "Hình 1.3: Caption lấy từ alt text")
    missing_inline.append(doc_pr)
    missing_drawing.append(missing_inline)
    missing_run.append(missing_drawing)
    missing_caption_drawing._p.append(missing_run)
    doc.add_paragraph("Trong đó:")
    doc.add_paragraph("$M$: Số hàng điểm ảnh")
    doc.add_paragraph("$N$: Số cột điểm ảnh")
    doc.save(path)

    assert publisher.finalize_word_docx(
        path,
        "Mục lục",
        page_start_heading="Lời nói đầu",
    ) is True

    with ZipFile(path) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        settings_xml = zf.read("word/settings.xml").decode("utf-8")
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
        footer_names = [name for name in zf.namelist() if name.startswith("word/footer")]
        assert footer_names
        footer_xml = zf.read(footer_names[0]).decode("utf-8")

    assert "Mục lục" in document_xml
    assert "Table of Contents" not in document_xml
    assert publisher._WORD_FIGURE_LIST_MARKER not in document_xml
    assert document_xml.count("<w:sectPr") >= 3
    assert 'w:vAlign w:val="center"' in document_xml
    assert 'w:left="1417"' in document_xml
    assert 'w:right="1417"' in document_xml
    assert 'w:top="1134"' in document_xml
    assert 'w:bottom="1134"' in document_xml
    assert 'w:start="1"' in document_xml
    assert "w:updateFields" in settings_xml
    assert "PAGE" in footer_xml
    assert 'w:jc w:val="center"' in footer_xml
    assert "PAGEREF fig_1" in document_xml
    assert "PAGEREF fig_3" in document_xml
    assert "HYPERLINK" in document_xml
    assert 'w:leader="dot"' in document_xml
    assert 'w:name="fig_1"' in document_xml
    assert 'w:name="fig_3"' in document_xml
    title_idx = document_xml.index("Giáo trình kiểm thử")
    toc_idx = document_xml.index("Mục lục")
    figure_title_idx = document_xml.index("Danh mục hình")
    preface_idx = document_xml.index("Lời nói đầu")
    assert title_idx < toc_idx < figure_title_idx < preface_idx
    assert (
        'w:type w:val="nextPage"' in document_xml[:toc_idx]
        or 'w:br w:type="page"' in document_xml[title_idx:toc_idx]
    )
    assert 'w:br w:type="page"' in document_xml[toc_idx:figure_title_idx]
    assert 'w:br w:type="page"' in document_xml[figure_title_idx:preface_idx]
    assert document_xml[figure_title_idx:preface_idx].count('w:type="page"') == 1

    figure_list_title = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?Danh mục hình(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert figure_list_title
    assert 'w:pStyle w:val="Heading1"' not in figure_list_title.group(0)
    assert 'w:jc w:val="center"' in figure_list_title.group(0)

    body_paragraph = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?Nội dung lời nói đầu cần được căn đều như bản PDF'
        r'(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert body_paragraph
    assert 'w:jc w:val="both"' in body_paragraph.group(0)

    heading_paragraph = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?CHƯƠNG 1: MỞ ĐẦU(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert heading_paragraph
    assert 'w:jc w:val="both"' not in heading_paragraph.group(0)
    assert 'w:jc w:val="center"' in heading_paragraph.group(0)

    code_paragraph = re.search(
        r"<w:p\b(?:(?!</w:p>).)*?left aligned"
        r'(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert code_paragraph
    assert 'w:jc w:val="left"' in code_paragraph.group(0)

    drawing_paragraph_xml = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?<w:drawing/?>(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert drawing_paragraph_xml
    assert 'w:jc w:val="center"' in drawing_paragraph_xml.group(0)

    caption_paragraphs = re.findall(
        r'<w:p\b(?:(?!</w:p>).)*?Hình 1.1: Sơ đồ kiểm thử'
        r'(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert caption_paragraphs
    assert any('w:jc w:val="center"' in paragraph for paragraph in caption_paragraphs)

    pattern_caption_paragraphs = re.findall(
        r'<w:p\b(?:(?!</w:p>).)*?Hình 1.2: Caption không có style riêng'
        r'(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert pattern_caption_paragraphs
    assert any('w:jc w:val="center"' in paragraph for paragraph in pattern_caption_paragraphs)

    inserted_caption_paragraphs = re.findall(
        r'<w:p\b(?:(?!</w:p>).)*?Hình 1.3: Caption lấy từ alt text'
        r'(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert inserted_caption_paragraphs
    assert any(
        'w:jc w:val="center"' in paragraph and "fldChar" not in paragraph
        for paragraph in inserted_caption_paragraphs
    )

    formula_intro = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?Trong đó:(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert formula_intro
    assert "w:keepNext" in formula_intro.group(0)
    assert "w:keepLines" in formula_intro.group(0)

    formula_item = re.search(
        r'<w:p\b(?:(?!</w:p>).)*?Số hàng điểm ảnh(?:(?!</w:p>).)*?</w:p>',
        document_xml,
    )
    assert formula_item
    assert "w:keepNext" in formula_item.group(0)
    assert "w:keepLines" in formula_item.group(0)

    for style_id in ("Heading1", "Heading2", "Heading3"):
        match = re.search(
            rf'<w:style\b[^>]*w:styleId="{style_id}"(?:(?!</w:style>).)*?</w:style>',
            styles_xml,
        )
        assert match, f"{style_id} style missing"
        style_xml = match.group(0)
        assert 'w:ascii="Times New Roman"' in style_xml
        assert "asciiTheme" not in style_xml


def test_reviewer_em_dash_cleanup_gate_is_vietnamese_only() -> None:
    assert ReviewerAgent._needs_em_dash_cleanup("CPU — bộ xử lý", "vi") is True
    assert ReviewerAgent._needs_em_dash_cleanup("CPU - bộ xử lý", "vi") is False
    assert ReviewerAgent._needs_em_dash_cleanup("CPU — processor", "en") is False


def test_remove_markdown_horizontal_rules_preserves_code_examples() -> None:
    content = (
        "## 1.2 Tiêu đề\n\n"
        "Đoạn một.\n\n"
        "---\n\n"
        "### 1.2.1 Mục con\n\n"
        "```markdown\n"
        "---\n"
        "title: Example\n"
        "---\n"
        "```\n\n"
        "----\n\n"
        "Đoạn hai."
    )

    cleaned = publisher.remove_markdown_horizontal_rules(content)

    assert "\n---\n\n###" not in cleaned
    assert "\n----\n\nĐoạn" not in cleaned
    assert "```markdown\n---\ntitle: Example\n---\n```" in cleaned


@pytest.mark.skipif(shutil.which("typst") is None, reason="Typst is not installed")
def test_special_title_compiles_with_typst(tmp_path: Path) -> None:
    block = publisher._build_title_block('Lập trình C# [Cơ bản] "2026" C:\\Temp')
    typst_source = block.removeprefix("```{=typst}\n").removesuffix("\n```")
    result = subprocess.run(
        ["typst", "compile", "--format", "pdf", "-", str(tmp_path / "title.pdf")],
        input=typst_source,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert (tmp_path / "title.pdf").stat().st_size > 0


def test_workflow_fails_when_a_required_export_is_missing(
    tmp_path: Path,
    monkeypatch,
) -> None:
    markdown = tmp_path / "book.md"
    docx = tmp_path / "book.docx"
    markdown.write_text("# Book", encoding="utf-8")
    docx.write_bytes(b"PK\x03\x04DOCX")

    class Workflow:
        def stream(self, state, config):
            yield {"publisher": {
                "final_markdown_filepath": str(markdown),
                "final_pdf_filepath": None,
                "final_docx_filepath": str(docx),
                "export_errors": {"pdf": "synthetic PDF failure"},
            }}

    monkeypatch.setattr(
        orchestrator,
        "create_content_after_confirm_workflow",
        lambda: Workflow(),
    )
    result = asyncio.run(workflow_runner.continue_after_curriculum_confirmation(
        textbook_id=1,
        confirmed_curriculum={
            "topic": "C#",
            "chapters": [{
                "title": "Chapter 1",
                "subsections": [{
                    "title": "Section 1",
                    "description": "Description",
                    "search_query": "C# basics",
                    "section_type": "medium",
                }],
            }],
        },
        initial_state={
            "request": "C#",
            "core_topic": "C#",
            "user_requirements": "",
            "language": "vi",
            "textbook_title": "C# Basics",
            "export_formats": ["PDF", "Word"],
        }, #type: ignore
        db=None,
    ))

    assert result["success"] is False
    assert "PDF" in result["error"]
    assert result["pdf_path"] is None
    assert result["docx_path"] == str(docx)


def test_workflow_reports_terminal_insufficient_context_before_export(
    monkeypatch,
) -> None:
    class Workflow:
        def stream(self, state, config):
            yield {"context_evaluator": {
                "context_quality": "insufficient",
                "current_chapter_index": 0,
                "current_subsection_index": 0,
                "rag_source_audit": {
                    "warnings": [
                        "Context did not meet strict RAG gate: chars=1293, chunks=1."
                    ],
                },
            }}

    monkeypatch.setattr(
        orchestrator,
        "create_content_after_confirm_workflow",
        lambda: Workflow(),
    )
    result = asyncio.run(workflow_runner.continue_after_curriculum_confirmation(
        textbook_id=1,
        confirmed_curriculum={
            "topic": "C#",
            "chapters": [{
                "title": "Chapter 1",
                "subsections": [{
                    "title": "Section 1",
                    "description": "Description",
                    "search_query": "C# basics",
                    "section_type": "medium",
                }],
            }],
        },
        initial_state={
            "request": "C#",
            "core_topic": "C#",
            "user_requirements": "",
            "language": "vi",
            "textbook_title": "C# Basics",
            "export_formats": ["PDF", "Word"],
        }, #type: ignore
        db=None,
    ))

    assert result["success"] is False
    assert "Insufficient RAG context for Chapter 1.1" in result["error"]
    assert "PDF" not in result["error"]
    assert result["source_audit"]["warnings"]
