"""Utilities for user-uploaded textbook sources.

Uploaded files are treated as first-party source material: they are extracted
directly, then passed into the same chunk/embedding pipeline as crawled pages.
The original file can be removed after extraction because the manifest carries
the citation metadata needed by the publisher.
"""

from __future__ import annotations

import json
import re
import unicodedata
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from fastapi import HTTPException, UploadFile
from langchain_core.documents import Document

from app.config import settings
from app.utils.log_config import setup_logger

logger = setup_logger(name="UploadedSources", logfile="logs/agents.log")

ALLOWED_SOURCE_SUFFIXES = {".pdf", ".docx"}
MAX_SOURCE_FILE_BYTES = 25 * 1024 * 1024
SOURCE_UPLOAD_DIR = settings.BASE_DIR / "source_uploads"
UPLOADED_SOURCE_TEXT_LOG = settings.BASE_DIR / "logs" / "uploaded_source_text.log"
UPLOADED_SOURCE_TEXT_LOG_MAX_CHARS = 30_000
VI_DIACRITIC_RE = re.compile(
    r"[àáảãạăắằẳẵặâấầẩẫậđèéẻẽẹêếềểễệìíỉĩịòóỏõọôốồổỗộơớờởỡợ"
    r"ùúủũụưứừửữựỳýỷỹỵđ]",
    re.IGNORECASE,
)


def _safe_filename(name: str) -> str:
    base = Path(name or "source").name
    base = re.sub(r"[^\w.\- ]+", "_", base, flags=re.UNICODE).strip()
    base = re.sub(r"\s+", "_", base)
    return base[:180] or "source"


def _source_id() -> str:
    return uuid.uuid4().hex[:16]


async def save_uploaded_source_files(
    textbook_id: int,
    files: list[UploadFile] | None,
) -> list[dict[str, Any]]:
    """Persist uploaded sources under a textbook-specific temporary directory."""
    manifests: list[dict[str, Any]] = []
    if not files:
        return manifests

    upload_dir = SOURCE_UPLOAD_DIR / str(textbook_id)
    upload_dir.mkdir(parents=True, exist_ok=True)

    for file in files:
        filename = _safe_filename(file.filename or "source")
        suffix = Path(filename).suffix.lower()
        if suffix not in ALLOWED_SOURCE_SUFFIXES:
            raise HTTPException(
                status_code=400,
                detail="Chỉ hỗ trợ file nguồn .pdf hoặc .docx.",
            )

        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail=f"File nguồn rỗng: {filename}")
        if len(content) > MAX_SOURCE_FILE_BYTES:
            raise HTTPException(
                status_code=400,
                detail=f"File nguồn vượt quá 25MB: {filename}",
            )

        source_id = _source_id()
        stored_name = f"{source_id}_{filename}"
        stored_path = upload_dir / stored_name
        stored_path.write_bytes(content)
        manifests.append({
            "id": source_id,
            "kind": "user_file",
            "type": suffix.lstrip("."),
            "filename": filename,
            "stored_path": str(stored_path),
            "status": "pending",
            "chunk_count": 0,
            "document_count": 0,
            "warnings": [],
            "apa": {
                "title": Path(filename).stem.replace("_", " "),
                "author": "",
                "year": "",
                "publisher": "",
                "url": "",
            },
        })

    return manifests


def url_source_manifest(urls: Iterable[str]) -> list[dict[str, Any]]:
    """Build manifest rows for user-provided URLs."""
    manifests: list[dict[str, Any]] = []
    seen: set[str] = set()
    for raw_url in urls or []:
        url = str(raw_url or "").strip()
        if not url or url in seen:
            continue
        seen.add(url)
        source_id = _source_id()
        manifests.append({
            "id": source_id,
            "kind": "user_url",
            "type": "url",
            "url": url,
            "status": "pending",
            "chunk_count": 0,
            "document_count": 0,
            "warnings": [],
            "apa": {
                "title": _title_from_url(url),
                "author": "",
                "year": "",
                "publisher": _publisher_from_url(url),
                "url": url,
            },
        })
    return manifests


def merge_source_materials(
    existing: list[dict[str, Any]] | None,
    additions: list[dict[str, Any]] | None,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in [*(existing or []), *(additions or [])]:
        key = str(item.get("url") or item.get("stored_path") or item.get("id") or "")
        if not key or key in seen:
            continue
        seen.add(key)
        rows.append(dict(item))
    return rows


def _title_from_url(url: str) -> str:
    try:
        from urllib.parse import urlparse

        parsed = urlparse(url)
        path = parsed.path.strip("/").split("/")[-1]
        title = path or parsed.netloc
        title = re.sub(r"[-_]+", " ", title)
        title = re.sub(r"\.(html?|pdf)$", "", title, flags=re.IGNORECASE)
        return title.strip() or url
    except Exception:
        return url


def _publisher_from_url(url: str) -> str:
    try:
        from urllib.parse import urlparse

        host = urlparse(url).netloc.lower().removeprefix("www.")
        parts = host.split(".")
        return parts[-2].replace("-", " ").title() if len(parts) >= 2 else host
    except Exception:
        return ""


def _clean_year(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return str(value.year)
    match = re.search(r"\b(19|20)\d{2}\b", str(value))
    return match.group(0) if match else ""


def _normalize_pdf_text_chars(value: str) -> str:
    return (
        str(value or "")
        .replace("ﬀ", "ff")
        .replace("ﬁ", "fi")
        .replace("ﬂ", "fl")
        .replace("ﬃ", "ffi")
        .replace("ﬄ", "ffl")
    )


def _clean_title(value: str, fallback: str) -> str:
    title = re.sub(r"\s+", " ", _normalize_pdf_text_chars(value)).strip()
    return title or fallback


def _ascii_key(value: str) -> str:
    normalized = unicodedata.normalize("NFD", str(value or ""))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", normalized).strip().lower()


def _looks_like_filename_title(value: str, filename: str) -> bool:
    title = _ascii_key(value)
    stem = _ascii_key(Path(filename or "").stem.replace("_", " "))
    if not title:
        return True
    if stem and (title == stem or title.replace(" ", "") == stem.replace(" ", "")):
        return True
    return bool(re.search(r"\.(pdf|docx?)$", title, re.IGNORECASE))


def _sentence_case_if_upper(value: str) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip(" :-–—\t\r\n"))
    if not text:
        return ""
    letters = [ch for ch in text if ch.isalpha()]
    if letters and sum(1 for ch in letters if ch.isupper()) / len(letters) >= 0.55:
        lowered = text.lower()
        text = lowered[:1].upper() + lowered[1:]
        roman_re = re.compile(r"\b(i{1,3}|iv|v|vi{0,3}|ix|x)\b", re.IGNORECASE)
        return roman_re.sub(lambda match: match.group(1).upper(), text)
    return text


def _normalize_work_title(value: str) -> str:
    text = _sentence_case_if_upper(value)
    match = re.match(r"^(bài\s*giảng|bai\s*giang|giáo\s*trình|giao\s*trinh)\s+(.+)$", text, re.IGNORECASE)
    if not match:
        return text
    label_key = _ascii_key(match.group(1))
    label = "Bài giảng" if label_key.startswith("bai giang") else "Giáo trình"
    rest = match.group(2).strip()
    rest = rest.lower()
    roman_re = re.compile(r"\b(i{1,3}|iv|v|vi{0,3}|ix|x)\b", re.IGNORECASE)
    rest = roman_re.sub(lambda roman_match: roman_match.group(1).upper(), rest)
    return f"{label} {rest}".strip()


def _looks_like_mojibake(value: str) -> bool:
    text = str(value or "")
    return bool(re.search(r"[ÃÏÐÑÕ]", text))


def _strip_academic_titles(value: str) -> str:
    return re.sub(
        r"\b(?:ths|thạc\s*sĩ|ts|tiến\s*sĩ|pgs\.?\s*ts|pgs|gs|ngưt|ngut|ks|cn)\.?\s*",
        "",
        str(value or ""),
        flags=re.IGNORECASE,
    )


def _proper_case_person_name(value: str) -> str:
    text = _strip_academic_titles(_normalize_pdf_text_chars(value))
    text = re.sub(r"\s+", " ", text).strip(" .;:-")
    if not text:
        return ""
    if _looks_like_mojibake(text):
        return text
    letters = [ch for ch in text if ch.isalpha()]
    if letters and sum(1 for ch in letters if ch.isupper()) / len(letters) >= 0.55:
        return text.lower().title()
    return text


def _clean_collective_author(value: str) -> str:
    text = re.sub(r"\s+", " ", _normalize_pdf_text_chars(value)).strip(" .;:-")
    if not text:
        return ""
    if _looks_like_mojibake(text):
        return text
    key = _ascii_key(text)
    if key.startswith("tap the giang vien"):
        rest = text[len("Tập thể giảng viên"):].strip(" .;:-")
        return "Tập thể giảng viên" + (f" {rest}" if rest else "")
    if key.startswith("nhom tac gia"):
        rest = text[len("Nhóm tác giả"):].strip(" .;:-")
        return "Nhóm tác giả" + (f" {rest}" if rest else "")
    if key.startswith("tap the tac gia"):
        rest = text[len("Tập thể tác giả"):].strip(" .;:-")
        return "Tập thể tác giả" + (f" {rest}" if rest else "")
    return _sentence_case_if_upper(text)


def _looks_like_author_line(line: str) -> bool:
    text = re.sub(r"\s+", " ", str(line or "")).strip(" .;:-")
    key = _ascii_key(text)
    if not text or len(text) > 140:
        return False
    if key.startswith((
        "loi noi dau", "muc luc", "chuong", "giao trinh", "bai giang",
        "nha xuat ban", "hieu chinh", "dung cho sinh vien", "thang ",
        "ha noi",
    )):
        return False
    if re.search(r"\b(ths|thạc sĩ|ts|tiến sĩ|pgs|gs|ks)\.?\b", key, re.IGNORECASE):
        return True
    words = [word for word in re.split(r"\s+", text) if word]
    if 2 <= len(words) <= 7:
        letters = [ch for ch in text if ch.isalpha()]
        upper_ratio = (
            sum(1 for ch in letters if ch.isupper()) / len(letters)
            if letters else 0.0
        )
        return upper_ratio >= 0.55 and not _is_cover_boilerplate(text)
    return False


def _normalize_author_entries(entries: list[str]) -> str:
    cleaned: list[str] = []
    seen: set[str] = set()
    for entry in entries:
        for part in re.split(r"\s*(?:;|,|\bvà\b|\band\b|[-–—])\s*", entry, flags=re.IGNORECASE):
            name = _proper_case_person_name(part)
            if not name:
                continue
            key = _ascii_key(name)
            if key and key not in seen:
                seen.add(key)
                cleaned.append(name)
    return ", ".join(cleaned)


def _is_cover_boilerplate(line: str) -> bool:
    key = _ascii_key(line)
    if not key or len(key) < 3:
        return True
    boilerplate = (
        "bo giao duc", "bo lao dong", "cong hoa", "doc lap", "tu do",
        "bo thong tin", "hoc vien ", "truong ", "khoa ", "phong ", "uy ban",
        "so giao duc", "loi noi dau", "muc luc", "tai lieu tham khao",
        "chuong ", "hinh ", "bang ", "isbn", "ma so", "luu hanh noi bo",
        "ptit", "bai giang", "giao trinh", "thang ", "ha noi", "hieu chinh",
        "dung cho sinh vien",
    )
    return key.startswith(boilerplate)


def _extract_collective_author_from_lines(lines: list[str]) -> str:
    for idx, line in enumerate(lines[:180]):
        key = _ascii_key(line)
        if not key.startswith(("tap the giang vien", "nhom tac gia", "tap the tac gia")):
            continue

        parts = [line.strip()]
        for next_line in lines[idx + 1: idx + 5]:
            next_key = _ascii_key(next_line)
            if not next_key:
                continue
            if next_key.startswith(("to mon", "khoa ", "bo mon", "nhom ", "tap the")):
                parts.append(next_line.strip())
                continue
            break
        return _clean_collective_author(" - ".join(parts))
    return ""


def _extract_authors_from_preface_lines(lines: list[str]) -> str:
    trigger_re = re.compile(
        r"(?:biên\s*soạn|bien\s*soan).{0,80}(?:nhóm\s*tác\s*giả|nhom\s*tac\s*gia|"
        r"tập\s*thể\s*tác\s*giả|tap\s*the\s*tac\s*gia|bởi|boi)",
        re.IGNORECASE,
    )
    inline_re = re.compile(
        r"(?:biên\s*soạn|bien\s*soan).{0,80}(?:bởi|boi)\s+(?P<authors>.+)$",
        re.IGNORECASE,
    )
    for idx, line in enumerate(lines[:180]):
        if not trigger_re.search(line):
            continue

        entries: list[str] = []
        inline_match = inline_re.search(line)
        if inline_match:
            inline_authors = re.split(
                r"\s+(?:theo|gồm|gom|để|de|phục vụ|phuc vu)\b",
                inline_match.group("authors").strip(" .;:-"),
                maxsplit=1,
                flags=re.IGNORECASE,
            )[0]
            if inline_authors and not _ascii_key(inline_authors).startswith(("nhom tac gia", "tap the tac gia")):
                entries.append(inline_authors)

        for next_line in lines[idx + 1: idx + 45]:
            key = _ascii_key(next_line)
            if not key:
                continue
            if key.startswith((
                "chuong ", "noi dung ", "trong qua trinh", "de bai giang",
                "xin chan thanh", "ha noi", "thang ", "tai lieu tham khao",
            )):
                if entries and not key.startswith("chuong "):
                    break
                continue
            if _looks_like_author_line(next_line):
                entries.append(next_line)

        if entries:
            return _normalize_author_entries(entries)
    return ""


def _extract_cover_author_before_title_marker(lines: list[str]) -> str:
    for idx, line in enumerate(lines[:80]):
        key = _ascii_key(line)
        if key not in {"bai giang", "giao trinh"}:
            continue

        entries: list[str] = []
        for prev_line in reversed(lines[max(0, idx - 10):idx]):
            prev_key = _ascii_key(prev_line)
            if not prev_key or prev_key in {"ptit"}:
                continue
            if _looks_like_author_line(prev_line):
                entries.append(prev_line)
                continue
            if entries:
                break

        if entries:
            return _normalize_author_entries(list(reversed(entries)))
    return ""


def _extract_author_from_cover_lines(lines: list[str]) -> str:
    label_re = re.compile(
        r"^(?:chủ\s*biên|chu\s*bien|biên\s*soạn|bien\s*soan|tác\s*giả|tac\s*gia|"
        r"người\s*biên\s*soạn|nguoi\s*bien\s*soan|nhóm\s*biên\s*soạn|nhom\s*bien\s*soan)"
        r"\s*[:\-–]\s*(.+)$",
        re.IGNORECASE,
    )
    bare_label_re = re.compile(
        r"^(?:chủ\s*biên|chu\s*bien|biên\s*soạn|bien\s*soan|tác\s*giả|tac\s*gia|"
        r"người\s*biên\s*soạn|nguoi\s*bien\s*soan|nhóm\s*biên\s*soạn|nhom\s*bien\s*soan)\s*$",
        re.IGNORECASE,
    )
    for idx, line in enumerate(lines[:120]):
        match = label_re.match(line.strip())
        if match:
            entries = [match.group(1)]
            for next_line in lines[idx + 1: idx + 6]:
                if _looks_like_author_line(next_line):
                    entries.append(next_line)
                    continue
                break
            return _normalize_author_entries(entries)
        if bare_label_re.match(_ascii_key(line)) and idx + 1 < len(lines):
            entries: list[str] = []
            for next_line in lines[idx + 1: idx + 7]:
                if _looks_like_author_line(next_line):
                    entries.append(next_line)
                    continue
                if entries:
                    break
            if entries:
                return _normalize_author_entries(entries)
    cover_author = _extract_cover_author_before_title_marker(lines)
    if cover_author:
        return cover_author
    preface_author = _extract_authors_from_preface_lines(lines)
    if preface_author:
        return preface_author
    return _extract_collective_author_from_lines(lines)


def _extract_publisher_from_cover_lines(lines: list[str]) -> str:
    english_publishers = (
        "No Starch Press, Inc.",
        "No Starch Press",
        "Pearson Education, Inc.",
        "Pearson",
        "Addison-Wesley",
        "O'Reilly Media",
        "O'Reilly",
        "Packt Publishing",
        "Manning Publications",
        "Wiley",
        "Springer",
        "CRC Press",
        "MIT Press",
    )
    for line in lines[:180]:
        text = _normalize_pdf_text_chars(line)
        for publisher in english_publishers:
            if re.search(rf"\b{re.escape(publisher)}\b", text, re.IGNORECASE):
                return publisher

    publisher_pattern_groups = (
        (r"^nhà\s*xuất\s*bản\b.+", r"^nha\s*xuat\s*ban\b.+"),
        (r"^học\s*viện\b.+", r"^hoc\s*vien\b.+"),
        (
            r"^trường\b.+",
            r"^truong\b.+",
            r"^đại\s*học\b.+",
            r"^dai\s*hoc\b.+",
            r"^cao\s*đẳng\b.+",
            r"^cao\s*dang\b.+",
        ),
        (r"^bộ\b.+", r"^bo\b.+"),
    )
    for patterns in publisher_pattern_groups:
        for line in lines[:80]:
            key = _ascii_key(line)
            if any(re.match(pattern, key, re.IGNORECASE) for pattern in patterns):
                cleaned = re.sub(r"\s+", " ", line).strip(" .;")
                if 5 <= len(cleaned) <= 140:
                    return cleaned
    return ""


def _title_case_english_title(value: str) -> str:
    text = re.sub(r"\s+", " ", _normalize_pdf_text_chars(value)).strip(" .;:-")
    if not text:
        return ""
    text = re.sub(r"\s+:\s+", ": ", text)
    titled = text.title()
    replacements = {
        "Wireshark": "Wireshark",
        "Tcp": "TCP",
        "Ip": "IP",
        "Dns": "DNS",
        "Http": "HTTP",
        "Https": "HTTPS",
        "Tls": "TLS",
        "Ssl": "SSL",
        "Ids": "IDS",
        "Idps": "IDPS",
        "Cyber": "Cyber",
    }
    for wrong, right in replacements.items():
        titled = re.sub(rf"\b{wrong}\b", right, titled)
    minor_words = "A|An|And|As|At|By|For|From|In|Of|On|Or|The|Through|To|With"
    titled = re.sub(
        rf"(?<!^)(?<!:\s)\b({minor_words})\b",
        lambda match: match.group(1).lower(),
        titled,
    )
    return titled


def _normalize_english_edition(value: str) -> str:
    key = _ascii_key(value)
    word_map = {
        "first": "1st",
        "second": "2nd",
        "third": "3rd",
        "fourth": "4th",
        "fifth": "5th",
        "sixth": "6th",
        "seventh": "7th",
        "eighth": "8th",
        "ninth": "9th",
        "tenth": "10th",
    }
    match = re.search(
        r"\b(first|second|third|fourth|fifth|sixth|seventh|eighth|ninth|tenth|\d+(?:st|nd|rd|th))\s+edition\b",
        key,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    edition = word_map.get(match.group(1).lower(), match.group(1).lower())
    return f"{edition} ed."


def _extract_english_book_metadata(lines: list[str]) -> dict[str, str]:
    metadata = {"title": "", "author": "", "publisher": "", "edition": ""}

    for line in lines[:180]:
        edition = _normalize_english_edition(line)
        if edition:
            metadata["edition"] = edition
            break

    for line in lines[:180]:
        text = _normalize_pdf_text_chars(line).strip()
        if " / " not in text:
            continue
        title_part, author_part = text.split(" / ", 1)
        if not re.search(r"[A-Za-z]{3,}", title_part) or not re.search(r"[A-Za-z]{3,}", author_part):
            continue
        if _ascii_key(title_part).startswith(("library of congress", "includes bibliographical")):
            continue
        title_part = re.sub(r"\s+:\s+", ": ", title_part)
        author_part = re.sub(r"\s*\b(?:p\.?\s*cm\.?|includes bibliographical references.*)$", "", author_part, flags=re.IGNORECASE)
        author_part = author_part.strip(" .;:-")
        author = _normalize_author_entries([author_part])
        title = _title_case_english_title(title_part)
        if title and author:
            metadata["title"] = title
            metadata["author"] = author
            break

    if not metadata["title"]:
        for idx, line in enumerate(lines[:80]):
            key = _ascii_key(line)
            if key in {"practical packet analysis", "network forensics"}:
                title_parts = [line]
                for next_line in lines[idx + 1: idx + 4]:
                    next_key = _ascii_key(next_line)
                    if not next_key:
                        continue
                    if next_key.startswith(("using ", "tracking ")):
                        title_parts.append(next_line)
                    break
                metadata["title"] = _title_case_english_title(": ".join(title_parts))
                break

    if not metadata["author"]:
        for idx, line in enumerate(lines[:100]):
            match = re.match(r"^by\s+(.+)$", _normalize_pdf_text_chars(line).strip(), flags=re.IGNORECASE)
            if match:
                metadata["author"] = _normalize_author_entries([match.group(1)])
                break
            if metadata["title"] and _ascii_key(line) in {
                _ascii_key(part)
                for part in re.split(r":\s*", metadata["title"])
            }:
                author_lines: list[str] = []
                for next_line in lines[idx + 1: idx + 6]:
                    next_key = _ascii_key(next_line)
                    if not next_key:
                        continue
                    if next_key.startswith(("upper saddle", "san francisco", "new york", "boston")):
                        break
                    if _looks_like_author_line(next_line):
                        author_lines.append(next_line)
                if author_lines:
                    metadata["author"] = _normalize_author_entries(author_lines)
                    break

    metadata["publisher"] = _extract_publisher_from_cover_lines(lines)
    return metadata


def _extract_year_from_cover_lines(lines: list[str]) -> str:
    date_line_re = re.compile(
        r"(?:hà\s*nội|ha\s*noi|tp\.?\s*hồ\s*chí\s*minh|tp\.?\s*ho\s*chi\s*minh|"
        r"tháng|thang).{0,40}\b((?:19|20)\d{2})\b",
        re.IGNORECASE,
    )
    for line in lines[:180]:
        key = _ascii_key(line)
        if not (
            key.startswith(("ha noi", "tp ho chi minh", "thang "))
            or (len(key) <= 60 and (" thang " in f" {key} " or key.startswith("thang ")))
        ):
            continue
        match = date_line_re.search(line)
        if match:
            return match.group(1)

    for line in lines[:180]:
        match = re.search(r"\bcopyright\b.{0,30}(?:©|\(c\))?\s*((?:19|20)\d{2})\b", line, re.IGNORECASE)
        if match:
            return match.group(1)

    years = re.findall(r"\b(19|20)\d{2}\b", "\n".join(lines[:120]))
    if years:
        full_years = re.findall(r"\b(?:19|20)\d{2}\b", "\n".join(lines[:120]))
        return full_years[-1]
    return ""


def _extract_title_from_preface_lines(lines: list[str]) -> str:
    for line in lines[:80]:
        stripped = line.strip(" .;:-")
        match = re.search(
            r"\b((?:bài\s*giảng|bai\s*giang|giáo\s*trình|giao\s*trinh)\s+.{4,90}?)"
            r"\s+(?:được|duoc|gồm|gom|cung\s*cấp|cung\s*cap|dùng|dung|biên\s*soạn|bien\s*soan)\b",
            stripped,
            flags=re.IGNORECASE,
        )
        if match:
            title = re.sub(r"\s+", " ", match.group(1)).strip(" .;:-")
            return _normalize_work_title(title)
    return ""


def _extract_title_from_cover_lines(lines: list[str], filename: str) -> str:
    lecture_label_re = re.compile(r"^(?:bài\s*giảng|bai\s*giang)\s*$", re.IGNORECASE)
    textbook_label_re = re.compile(r"^(?:giáo\s*trình|giao\s*trinh)\s*$", re.IGNORECASE)
    inline_work_re = re.compile(
        r"^(?P<label>bài\s*giảng|bai\s*giang|giáo\s*trình|giao\s*trinh)"
        r"\s*[:\-–]?\s*(?P<title>.+)$",
        re.IGNORECASE,
    )

    def collect_title_after(idx: int, prefix: str) -> str:
        for next_line in lines[idx + 1: idx + 6]:
            stripped_next = next_line.strip(" .;:-")
            key_next = _ascii_key(stripped_next)
            if not stripped_next:
                continue
            if key_next.startswith((
                "thang ", "ha noi", "tp ho chi minh", "dung cho sinh vien",
                "nghe ", "mon hoc", "hoc phan", "chu bien", "bien soan",
                "tac gia",
            )):
                break
            if stripped_next.startswith("("):
                continue
            if _is_cover_boilerplate(stripped_next):
                continue
            if 4 <= len(stripped_next) <= 120:
                return _normalize_work_title(f"{prefix} {stripped_next}")
        return ""

    for idx, line in enumerate(lines[:100]):
        stripped = line.strip()
        key = _ascii_key(stripped)
        if not key:
            continue
        match = re.search(
            r"(?:môn\s*học|môn\s*học/mô\s*đun|mon\s*hoc|mon\s*hoc/mo\s*dun|học\s*phần|hoc\s*phan)"
            r"\s*[:\-–]\s*(.+)$",
            stripped,
            re.IGNORECASE,
        )
        if match:
            return _sentence_case_if_upper(match.group(1))
        match = re.search(r"giáo\s*trình\s*[:\-–]\s*(.+)$", stripped, re.IGNORECASE)
        if match:
            return _normalize_work_title(f"Giáo trình {match.group(1)}")
        match = inline_work_re.match(stripped)
        if match and match.group("title").strip():
            if re.search(
                r"\s(?:được|duoc|gồm|gom|cung\s*cấp|cung\s*cap|dùng|dung|biên\s*soạn|bien\s*soan)\b",
                stripped,
                flags=re.IGNORECASE,
            ):
                continue
            label = "Bài giảng" if _ascii_key(match.group("label")).startswith("bai giang") else "Giáo trình"
            return _normalize_work_title(f"{label} {match.group('title')}")
        if lecture_label_re.match(stripped):
            title = collect_title_after(idx, "Bài giảng")
            if title:
                return title
        if key == "giao trinh":
            title = collect_title_after(idx, "Giáo trình")
            if title:
                return title

    preface_title = _extract_title_from_preface_lines(lines)
    if preface_title:
        return preface_title

    candidates: list[str] = []
    for line in lines[:100]:
        stripped = line.strip(" .;:-")
        if not (4 <= len(stripped) <= 130):
            continue
        if _is_cover_boilerplate(stripped):
            continue
        if _looks_like_filename_title(stripped, filename):
            continue
        letters = [ch for ch in stripped if ch.isalpha()]
        upper_ratio = (
            sum(1 for ch in letters if ch.isupper()) / len(letters)
            if letters else 0.0
        )
        if upper_ratio >= 0.55 or re.search(r"\b(giáo trình|giao trinh|kỹ thuật|ky thuat|điện|dien)\b", stripped, re.IGNORECASE):
            candidates.append(stripped)
    return _sentence_case_if_upper(candidates[0]) if candidates else ""


def infer_apa_metadata_from_text(
    text: str,
    *,
    filename: str = "",
    base: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Infer APA metadata from the first pages of Vietnamese textbook files."""
    base = dict(base or {})
    lines = [
        re.sub(r"\s+", " ", _normalize_pdf_text_chars(line)).strip()
        for line in str(text or "").splitlines()
        if re.sub(r"\s+", " ", line).strip()
    ][:180]
    english_metadata = _extract_english_book_metadata(lines)

    warnings: list[str] = []
    inferred = {
        "title": english_metadata.get("title") or _extract_title_from_cover_lines(lines, filename),
        "author": english_metadata.get("author") or _extract_author_from_cover_lines(lines),
        "year": _extract_year_from_cover_lines(lines),
        "publisher": english_metadata.get("publisher") or _extract_publisher_from_cover_lines(lines),
        "edition": english_metadata.get("edition") or str(base.get("edition") or ""),
        "url": "",
    }
    for field in ("title", "author", "publisher"):
        if inferred.get(field) and _looks_like_mojibake(str(inferred[field])):
            warnings.append(
                "Metadata APA có dấu hiệu lỗi font tiếng Việt; vui lòng kiểm tra/cập nhật thủ công."
            )
            if field == "author":
                inferred[field] = ""

    current_title = str(base.get("title") or "")
    if not inferred["title"] and not _looks_like_filename_title(current_title, filename):
        inferred["title"] = current_title

    if not inferred["author"]:
        warnings.append(
            "Không xác định chắc chắn tác giả hoặc tập thể tác giả từ nội dung nguồn; "
            "vui lòng kiểm tra metadata APA."
        )
    if not inferred["year"]:
        warnings.append("Không xác định được năm xuất bản từ nội dung nguồn.")
    if not inferred["publisher"]:
        warnings.append("Không xác định được đơn vị xuất bản từ nội dung nguồn.")

    result = {
        "title": inferred["title"] or str(base.get("title") or Path(filename).stem.replace("_", " ")),
        "author": inferred["author"],
        "year": inferred["year"] or str(base.get("year") or ""),
        "publisher": inferred["publisher"] or str(base.get("publisher") or ""),
        "edition": inferred["edition"],
        "url": str(base.get("url") or ""),
        "needs_review": bool(warnings),
        "confidence": "needs_review" if warnings else "high",
        "warnings": list(dict.fromkeys(warnings)),
    }
    return result


def detect_source_language(text: str) -> str:
    """Fast VI/EN heuristic used for uploaded source manifests."""
    sample = str(text or "")[:5000]
    return "vi" if VI_DIACRITIC_RE.search(sample) else "en"


def user_source_language_profile(source_materials: list[dict[str, Any]] | None) -> dict[str, Any]:
    rows = [
        item for item in source_materials or []
        if item.get("kind") in {"user_file", "user_url"}
    ]
    language_counts: Counter[str] = Counter(
        str(item.get("language") or "").lower()
        for item in rows
        if item.get("language")
    )
    file_rows = [item for item in rows if item.get("kind") == "user_file"]
    vi_count = int(language_counts.get("vi", 0))
    en_count = int(language_counts.get("en", 0))
    total_known = vi_count + en_count
    preference = ""
    if total_known and vi_count >= en_count and vi_count / total_known >= 0.5:
        preference = "vi"
    elif total_known and en_count > vi_count and en_count / total_known >= 0.6:
        preference = "en"
    return {
        "preference": preference,
        "language_counts": dict(language_counts),
        "user_source_count": len(rows),
        "user_file_count": len(file_rows),
        "known_language_count": total_known,
        "vi_ratio": (vi_count / total_known) if total_known else 0.0,
    }


def _normalize_line(line: str) -> str:
    return re.sub(r"\s+", " ", str(line or "").strip()).lower()


def _remove_repeated_lines(pages: list[str]) -> list[str]:
    """Best-effort removal of text-layer watermarks/header/footer repeats."""
    normalized_counts: Counter[str] = Counter()
    page_lines: list[list[str]] = []
    for text in pages:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        page_lines.append(lines)
        normalized_counts.update(set(_normalize_line(line) for line in lines if line.strip()))

    page_count = max(1, len(pages))
    threshold = max(2, int(page_count * 0.35))
    repeated = {
        line
        for line, count in normalized_counts.items()
        if count >= threshold and (len(line) <= 140 or _looks_like_watermark(line))
    }
    if not repeated:
        return pages

    cleaned_pages: list[str] = []
    for lines in page_lines:
        kept = [
            line
            for line in lines
            if _normalize_line(line) not in repeated
        ]
        cleaned_pages.append("\n".join(kept))
    return cleaned_pages


def _looks_like_watermark(line: str) -> bool:
    return bool(re.search(
        r"\b(draft|confidential|sample|watermark|copy|internal|không sao chép|"
        r"ban nháp|nội bộ|mẫu)\b",
        line,
        re.IGNORECASE,
    ))


_CONTENT_START_RE = re.compile(
    r"^\s*(?:"
    r"(?:chương|chuong|chapter)\s+(?:\d+|[ivxlcdm]+)\b"
    r"|(?:bài|bai|unit|module|lesson)\s+\d+\b"
    r"|(?:\d{1,2})(?:\.\d{1,2}){0,3}\s+[A-Za-zÀ-ỹ]"
    r")",
    re.IGNORECASE,
)
_BACK_MATTER_HEADING_RE = re.compile(
    r"^\s*(?:"
    r"tài\s*liệu\s*tham\s*khảo|tai\s*lieu\s*tham\s*khao|"
    r"tham\s*khảo|tham\s*khao|references|bibliography|works\s+cited|"
    r"index|chỉ\s*mục|chi\s*muc"
    r")\s*$",
    re.IGNORECASE,
)
_FRONT_LIST_HEADING_RE = re.compile(
    r"^\s*(?:"
    r"mục\s*lục|muc\s*luc|contents|table\s+of\s+contents|"
    r"danh\s*mục\s*hình|danh\s*muc\s*hinh|danh\s*mục\s*bảng|danh\s*muc\s*bang|"
    r"list\s+of\s+figures|list\s+of\s+tables|list\s+of\s+illustrations|"
    r"list\s+of\s+examples|list\s+of\s+listings"
    r")\s*$",
    re.IGNORECASE,
)


def _looks_like_toc_listing_line(line: str) -> bool:
    stripped = line.strip()
    if re.match(r"^(?:chapter|chương|chuong)\s+\d+\s*,", stripped, re.IGNORECASE):
        return True
    if re.match(r"^(?:chapter|chương|chuong)\s+(?:\d+|[ivxlcdm]+)\s*$", stripped, re.IGNORECASE):
        return False
    if re.search(r"\.{3,}\s*\d{1,4}\s*$", stripped):
        return True
    if _CONTENT_START_RE.match(stripped) and re.search(r"\s{2,}\d{1,4}\s*$", stripped):
        return True
    if _CONTENT_START_RE.match(stripped) and len(stripped.split()) <= 8 and re.search(r"\b\d{1,4}\s*$", stripped):
        return True
    return False


def _is_front_listing_page(lines: list[str]) -> bool:
    non_empty = [line.strip() for line in lines if line.strip()]
    if not non_empty:
        return False
    if any(_FRONT_LIST_HEADING_RE.match(line.strip(" .:-–—\t")) for line in non_empty[:40]):
        return True
    if len(non_empty) >= 6:
        listing_count = sum(1 for line in non_empty if _looks_like_toc_listing_line(line))
        if listing_count >= 3 and listing_count / len(non_empty) >= 0.25:
            return True
    return False


def _find_content_start_line(lines: list[str]) -> int | None:
    if _is_front_listing_page(lines):
        return None
    for idx, line in enumerate(lines):
        stripped = line.strip(" .:-–—\t")
        if not stripped:
            continue
        if _BACK_MATTER_HEADING_RE.match(stripped):
            return None
        if _looks_like_toc_listing_line(line):
            continue
        if _CONTENT_START_RE.match(stripped):
            return idx
    return None


def _find_back_matter_line(lines: list[str]) -> int | None:
    if _is_front_listing_page(lines):
        return None
    for idx, line in enumerate(lines):
        stripped = line.strip(" .:-–—\t")
        if _BACK_MATTER_HEADING_RE.match(stripped):
            return idx
    return None


def _clean_source_pages_for_embedding(pages: list[str]) -> tuple[list[str], dict[str, int]]:
    """
    Remove front/back matter before chunking while keeping metadata extraction intact.

    Front matter is dropped only when a reliable content start marker is found,
    which prevents accidentally deleting whole documents with uncommon structure.
    """

    original_chars = sum(len(page or "") for page in pages)
    cleaned_pages = list(pages or [])
    stats = {
        "original_chars": original_chars,
        "cleaned_chars": original_chars,
        "removed_front_pages": 0,
        "removed_back_pages": 0,
        "removed_chars": 0,
    }
    if not cleaned_pages:
        return cleaned_pages, stats

    start_page: int | None = None
    start_line: int | None = None
    for page_idx, page in enumerate(cleaned_pages):
        lines = [line for line in str(page or "").splitlines()]
        found = _find_content_start_line(lines)
        if found is not None:
            start_page = page_idx
            start_line = found
            break

    if start_page is not None:
        stats["removed_front_pages"] = start_page
        cleaned_pages = cleaned_pages[start_page:]
        if cleaned_pages and start_line:
            first_lines = cleaned_pages[0].splitlines()
            cleaned_pages[0] = "\n".join(first_lines[start_line:])

    earliest_back_page = max(0, int(len(cleaned_pages) * 0.60))
    for page_idx, page in enumerate(cleaned_pages):
        if page_idx < earliest_back_page:
            continue
        lines = [line for line in str(page or "").splitlines()]
        found = _find_back_matter_line(lines)
        if found is None:
            continue
        before = "\n".join(lines[:found]).strip()
        tail_count = len(cleaned_pages) - page_idx
        stats["removed_back_pages"] = tail_count
        cleaned_pages = cleaned_pages[:page_idx]
        if before:
            cleaned_pages.append(before)
        break

    cleaned_pages = [page.strip() for page in cleaned_pages if str(page or "").strip()]
    cleaned_chars = sum(len(page) for page in cleaned_pages)
    stats["cleaned_chars"] = cleaned_chars
    stats["removed_chars"] = max(0, original_chars - cleaned_chars)

    if original_chars >= 1000 and cleaned_chars < max(500, int(original_chars * 0.10)):
        logger.warning(
            "Source text front/back matter cleanup removed too much content; keeping original text"
        )
        return pages, {
            **stats,
            "cleaned_chars": original_chars,
            "removed_front_pages": 0,
            "removed_back_pages": 0,
            "removed_chars": 0,
        }

    return cleaned_pages, stats


def clean_source_text_for_embedding(text: str) -> tuple[str, dict[str, int]]:
    """Line-based companion for DOCX or crawled PDF text without page metadata."""

    original = str(text or "")
    original_chars = len(original)
    stats = {
        "original_chars": original_chars,
        "cleaned_chars": original_chars,
        "removed_front_pages": 0,
        "removed_back_pages": 0,
        "removed_chars": 0,
    }
    lines = original.splitlines()
    if not lines:
        return original, stats

    start_idx: int | None = None
    in_front_listing = False
    for idx, line in enumerate(lines):
        stripped = line.strip(" .:-–—\t")
        if not stripped:
            continue
        if _FRONT_LIST_HEADING_RE.match(stripped):
            in_front_listing = True
            continue
        if in_front_listing and _looks_like_toc_listing_line(line):
            continue
        if _BACK_MATTER_HEADING_RE.match(stripped):
            break
        if _looks_like_toc_listing_line(line):
            continue
        if _CONTENT_START_RE.match(stripped):
            start_idx = idx
            break

    if start_idx is None:
        return original, stats

    kept_lines = lines[start_idx:]
    back_idx = _find_back_matter_line(kept_lines)
    if back_idx is not None and back_idx >= int(len(kept_lines) * 0.60):
        kept_lines = kept_lines[:back_idx]
        stats["removed_back_pages"] = 1

    cleaned = "\n".join(kept_lines).strip()
    cleaned_chars = len(cleaned)
    stats["cleaned_chars"] = cleaned_chars
    stats["removed_front_pages"] = 1 if start_idx > 0 else 0
    stats["removed_chars"] = max(0, original_chars - cleaned_chars)

    if original_chars >= 1000 and cleaned_chars < max(500, int(original_chars * 0.10)):
        logger.warning(
            "Source text front/back matter cleanup removed too much content; keeping original text"
        )
        return original, {
            **stats,
            "cleaned_chars": original_chars,
            "removed_front_pages": 0,
            "removed_back_pages": 0,
            "removed_chars": 0,
        }

    return cleaned, stats


def _log_uploaded_source_text(
    entry: dict[str, Any],
    text: str,
    *,
    page: int | None = None,
    stage: str = "extracted_cleaned",
) -> None:
    """Write cleaned uploaded-source text to a dedicated debug log."""

    try:
        cleaned_text = str(text or "").strip()
        if not cleaned_text:
            return

        UPLOADED_SOURCE_TEXT_LOG.parent.mkdir(parents=True, exist_ok=True)
        apa = entry.get("apa") or {}
        metadata = {
            "source_id": entry.get("id") or "",
            "filename": entry.get("filename") or "",
            "type": entry.get("type") or "",
            "page": page,
            "stage": stage,
            "chars": len(cleaned_text),
            "language": entry.get("language") or "",
            "apa": {
                "title": apa.get("title") or "",
                "author": apa.get("author") or "",
                "year": apa.get("year") or "",
                "publisher": apa.get("publisher") or "",
                "needs_review": apa.get("needs_review") or False,
                "warnings": apa.get("warnings") or [],
            },
        }
        visible_text = cleaned_text[:UPLOADED_SOURCE_TEXT_LOG_MAX_CHARS]
        if len(cleaned_text) > UPLOADED_SOURCE_TEXT_LOG_MAX_CHARS:
            visible_text += (
                f"\n\n[TRUNCATED: showing first {UPLOADED_SOURCE_TEXT_LOG_MAX_CHARS} "
                f"of {len(cleaned_text)} characters]"
            )

        with open(UPLOADED_SOURCE_TEXT_LOG, "a", encoding="utf-8") as handle:
            handle.write("\n" + "=" * 100 + "\n")
            handle.write("[UPLOADED SOURCE TEXT]\n")
            handle.write(json.dumps(metadata, ensure_ascii=False, indent=2))
            handle.write("\n--- TEXT START ---\n")
            handle.write(visible_text)
            handle.write("\n--- TEXT END ---\n")
    except Exception as exc:
        logger.warning("Could not write uploaded source text debug log: %s", exc)


def _extract_pdf(entry: dict[str, Any]) -> tuple[list[Document], dict[str, Any]]:
    path = Path(str(entry.get("stored_path") or ""))
    docs: list[Document] = []
    warnings = list(entry.get("warnings") or [])
    try:
        import fitz  # type: ignore

        pdf = fitz.open(path)
        meta = pdf.metadata or {}
        pages = [page.get_text("text") or "" for page in pdf]
        pages = _remove_repeated_lines(pages)
        entry["language"] = detect_source_language("\n".join(pages))
        title = _clean_title(meta.get("title", ""), Path(entry.get("filename", path.name)).stem)
        author = _clean_title(meta.get("author", ""), "")
        year = _clean_year(meta.get("creationDate") or meta.get("modDate"))
        base_apa = {
            **(entry.get("apa") or {}),
            "title": title,
            "author": author,
            "year": year,
            "publisher": entry.get("apa", {}).get("publisher", ""),
            "url": "",
        }
        entry["apa"] = infer_apa_metadata_from_text(
            "\n".join(pages[:10]),
            filename=str(entry.get("filename") or path.name),
            base=base_apa,
        )
        warnings.extend(
            warning
            for warning in entry["apa"].get("warnings", [])
            if warning not in warnings
        )
        entry["apa_needs_review"] = bool(entry["apa"].get("needs_review"))
        embedding_pages, filter_stats = _clean_source_pages_for_embedding(pages)
        entry["embedding_text_filter"] = filter_stats
        if filter_stats.get("removed_chars", 0) > 0:
            logger.info(
                "Uploaded PDF front/back matter removed for %s: %s",
                entry.get("filename") or path.name,
                filter_stats,
            )
        page_offset = int(filter_stats.get("removed_front_pages") or 0)
        for idx, text in enumerate(embedding_pages, start=1 + page_offset):
            if len(text.strip()) < 100:
                continue
            _log_uploaded_source_text(entry, text, page=idx)
            docs.append(_document_for_uploaded_source(entry, text, page=idx))
        pdf.close()
    except Exception as exc:
        warnings.append(f"Không extract được PDF: {exc}")
        logger.warning("PDF source extraction failed for %s: %s", path, exc)

    entry["warnings"] = warnings
    return docs, entry


def _extract_docx(entry: dict[str, Any]) -> tuple[list[Document], dict[str, Any]]:
    path = Path(str(entry.get("stored_path") or ""))
    warnings = list(entry.get("warnings") or [])
    try:
        from docx import Document as DocxDocument

        document = DocxDocument(path)  # type: ignore
        props = document.core_properties
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if text:
                lines.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [" ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    lines.append(row_text)

        text = "\n".join(_remove_repeated_lines(["\n".join(lines)]))
        entry["language"] = detect_source_language(text)
        title = _clean_title(getattr(props, "title", ""), Path(entry.get("filename", path.name)).stem)
        author = _clean_title(getattr(props, "author", ""), "")
        year = _clean_year(getattr(props, "created", None) or getattr(props, "modified", None))
        base_apa = {
            **(entry.get("apa") or {}),
            "title": title,
            "author": author,
            "year": year,
            "publisher": entry.get("apa", {}).get("publisher", ""),
            "url": "",
        }
        entry["apa"] = infer_apa_metadata_from_text(
            text[:10000],
            filename=str(entry.get("filename") or path.name),
            base=base_apa,
        )
        warnings.extend(
            warning
            for warning in entry["apa"].get("warnings", [])
            if warning not in warnings
        )
        entry["apa_needs_review"] = bool(entry["apa"].get("needs_review"))
        embedding_text, filter_stats = clean_source_text_for_embedding(text)
        entry["embedding_text_filter"] = filter_stats
        if filter_stats.get("removed_chars", 0) > 0:
            logger.info(
                "Uploaded DOCX front/back matter removed for %s: %s",
                entry.get("filename") or path.name,
                filter_stats,
            )
        if len(embedding_text.strip()) >= 100:
            _log_uploaded_source_text(entry, embedding_text, page=1)
        docs = [_document_for_uploaded_source(entry, embedding_text, page=1)] if len(embedding_text.strip()) >= 100 else []
    except Exception as exc:
        warnings.append(f"Không extract được DOCX: {exc}")
        logger.warning("DOCX source extraction failed for %s: %s", path, exc)
        docs = []

    entry["warnings"] = warnings
    return docs, entry


def _document_for_uploaded_source(entry: dict[str, Any], text: str, *, page: int) -> Document:
    source_id = str(entry.get("id") or "")
    filename = str(entry.get("filename") or "uploaded-source")
    apa = entry.get("apa") or {}
    source_url = f"uploaded://{source_id}/{filename}"
    return Document(
        page_content=text,
        metadata={
            "source": source_url,
            "source_url": source_url,
            "source_kind": "user_file",
            "user_source": "true",
            "direct_custom_url": "true",
            "source_id": source_id,
            "source_title": str(apa.get("title") or Path(filename).stem),
            "source_author": str(apa.get("author") or ""),
            "source_year": str(apa.get("year") or ""),
            "source_publisher": str(apa.get("publisher") or ""),
            "language": str(entry.get("language") or ""),
            "filename": filename,
            "type": str(entry.get("type") or "file"),
            "page": int(page),
            "domain": "uploaded-file",
            "trusted_source": "true",
        },
    )


def extract_uploaded_source_documents(
    source_materials: list[dict[str, Any]] | None,
    *,
    cleanup_files: bool = True,
) -> tuple[list[Document], list[dict[str, Any]]]:
    """Extract pending uploaded file sources into LangChain documents."""
    docs: list[Document] = []
    updated: list[dict[str, Any]] = []

    for raw_entry in source_materials or []:
        entry = dict(raw_entry)
        if entry.get("kind") != "user_file":
            updated.append(entry)
            continue

        source_type = str(entry.get("type") or "").lower()
        if source_type == "pdf":
            entry_docs, entry = _extract_pdf(entry)
        elif source_type == "docx":
            entry_docs, entry = _extract_docx(entry)
        else:
            entry_docs = []
            entry["warnings"] = [*(entry.get("warnings") or []), "Loại file không được hỗ trợ."]

        entry["document_count"] = len(entry_docs)
        entry["status"] = "ready" if entry_docs else "failed"
        docs.extend(entry_docs)

        if cleanup_files and entry_docs:
            try:
                path = Path(str(entry.get("stored_path") or ""))
                if path.exists():
                    path.unlink()
                entry["stored_path"] = ""
            except Exception as exc:
                entry["warnings"] = [*(entry.get("warnings") or []), f"Không xóa được file tạm: {exc}"]
        updated.append(entry)

    return docs, updated


def update_url_manifest_from_clean_links(
    source_materials: list[dict[str, Any]] | None,
    clean_links: list[dict[str, str]],
) -> list[dict[str, Any]]:
    """Mark user URL manifest rows as crawlable/failed after URL filtering."""
    by_url = {item.get("url"): item for item in clean_links}
    updated: list[dict[str, Any]] = []
    for raw_entry in source_materials or []:
        entry = dict(raw_entry)
        if entry.get("kind") != "user_url":
            updated.append(entry)
            continue
        clean = by_url.get(entry.get("url"))
        if clean:
            entry["status"] = "ready"
            entry["type"] = clean.get("type", "url")
            apa = dict(entry.get("apa") or {})
            if clean.get("search_title") and not apa.get("title"):
                apa["title"] = clean["search_title"]
            entry["apa"] = apa
        else:
            entry["status"] = "failed"
            entry["warnings"] = [*(entry.get("warnings") or []), "URL không vượt qua bước lọc/crawl."]
        updated.append(entry)
    return updated


def format_apa_numbered_references(
    source_materials: list[dict[str, Any]] | None,
    *,
    language: str = "vi",
) -> list[str]:
    """Return numbered APA-like references for valid user-provided sources."""
    rows: list[str] = []
    valid = [
        item for item in source_materials or []
        if item.get("kind") in {"user_url", "user_file"}
        and item.get("status") in {"ready", "embedded", "completed"}
    ]
    for index, item in enumerate(valid, start=1):
        apa = item.get("apa") or {}
        title = _clean_title(str(apa.get("title") or ""), str(item.get("filename") or item.get("url") or "Nguồn tài liệu"))
        author = str(apa.get("author") or "").strip()
        year = str(apa.get("year") or "n.d.").strip() or "n.d."
        publisher = str(apa.get("publisher") or "").strip()
        edition = str(apa.get("edition") or "").strip()
        url = str(apa.get("url") or item.get("url") or "").strip()

        lead = author or publisher or title
        parts = [f"[{index}] {lead}. ({year})."]
        if author or publisher:
            title_part = f"{title} ({edition})." if edition else title + "."
            parts.append(title_part)
        if publisher and _ascii_key(publisher) != _ascii_key(lead):
            parts.append(publisher + ".")
        if url:
            parts.append(url)
        rows.append(" ".join(parts))
    return rows
