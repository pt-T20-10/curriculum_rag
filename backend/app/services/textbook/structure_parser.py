"""Helpers for user-provided textbook structures."""

from __future__ import annotations

from io import BytesIO
import re
import unicodedata
from typing import Any


class StructureParseError(ValueError):
    """Raised when a user-provided outline cannot be parsed safely."""


_HEADING_RE = re.compile(r"^(#{2,4})\s+(.+?)\s*$")
_SUBSECTION_PREFIX_RE = re.compile(
    r"^(?:"
    r"\(?\d+(?:\.\d+)+\)?"
    r"|(?:mục|section)\s+\d+(?:\.\d+)+"
    r")\s*[:.)\-\s]*",
    re.IGNORECASE,
)
_CHAPTER_RE = re.compile(
    r"^(?:chương|chuong|bài|bai)(?:\s+thực\s+hành|\s+thuc\s+hanh)?\s*(\d+)\b",
    re.IGNORECASE,
)
_NUMBERED_ITEM_RE = re.compile(r"(?<![\d\[])(\d{1,2})\.(\d{1,2})(?:\.(\d{1,2}))?\.?\s+")
_TOPIC_RE = re.compile(r"t[êe]n\s+gi[aá]o\s+tr[iì]nh(?:/s[aá]ch)?\s*:\s*(.*)", re.IGNORECASE)
_MAX_STRUCTURE_FILE_SIZE = 10 * 1024 * 1024


def _normalize_ascii(value: Any) -> str:
    text = str(value or "").lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d")
    return re.sub(r"\s+", " ", text).strip()


def _strip_refs(value: str) -> str:
    text = re.sub(r"\s*\[[^\]]{1,20}\]", "", value)
    text = re.sub(r"(?:\s*[,;]?\s*\[?\d+\]?)+\s*$", "", text)
    return text.strip()


def _clean_outline_title(value: Any) -> str:
    text = _strip_refs(str(value or ""))
    text = re.sub(r"\s+", " ", text).strip(" \t\r\n:;,.")
    text = re.sub(
        r"\b(?:tóm tắt chương|tom tat chuong)\b.*$",
        "",
        text,
        flags=re.IGNORECASE,
    ).strip(" \t\r\n:;,.")
    practice_match = re.search(
        r"\b(?:bài\s*tập|bai\s*tap|câu\s+hỏi|cau\s+hoi)\b.*$",
        text,
        flags=re.IGNORECASE,
    )
    if practice_match and practice_match.start() > 0:
        text = text[:practice_match.start()].strip(" \t\r\n:;,.")
    return text


def _is_boilerplate(value: Any) -> bool:
    normalized = _normalize_ascii(value)
    if not normalized:
        return True
    boilerplate_exact = {
        "tom tat chuong",
        "bai tap",
        "bai tap chuong",
        "cau hoi on tap",
        "cau hoi trac nghiem on tap",
        "bai tap cau hoi on tap va huong dan tai lieu doc them neu co",
        "tong so trang",
    }
    if normalized in boilerplate_exact:
        return True
    if normalized.startswith("tom tat chuong"):
        return True
    generic_starts = (
        "bai tap chuong",
        "bai tap cau hoi",
        "cau hoi on tap",
        "cau hoi trac nghiem",
        "huong dan tai lieu doc them",
    )
    return normalized.startswith(generic_starts)


def _parse_page_count(value: Any) -> int | None:
    text = str(value or "").replace("~", " ")
    matches = re.findall(r"\d{1,4}", text)
    if not matches:
        return None
    try:
        pages = int(matches[-1])
    except ValueError:
        return None
    return pages if 1 <= pages <= 2000 else None


def _add_warning(warnings: list[str], message: str) -> None:
    if message not in warnings:
        warnings.append(message)


def _strip_chapter_prefix(text: str, chapter_number: int | None = None) -> str:
    if chapter_number is None:
        pattern = r"^\s*(?:chương|chuong|bài|bai)(?:\s+thực\s+hành|\s+thuc\s+hanh)?\s*\d+\s*[:.)\-]*\s*"
    else:
        pattern = rf"^\s*(?:chương|chuong|bài|bai)(?:\s+thực\s+hành|\s+thuc\s+hanh)?\s*{chapter_number}\s*[:.)\-]*\s*"
    return re.sub(pattern, "", text, flags=re.IGNORECASE).strip()


def _chapter_title_from_content(content: str, chapter_number: int | None = None) -> str:
    without_prefix = _strip_chapter_prefix(content, chapter_number)
    first_numbered = _NUMBERED_ITEM_RE.search(without_prefix)
    if first_numbered:
        without_prefix = without_prefix[:first_numbered.start()]
    return _clean_outline_title(without_prefix)


def _split_numbered_items(text: str) -> list[dict[str, Any]]:
    content = re.sub(r"\s+", " ", str(text or "")).strip()
    matches = list(_NUMBERED_ITEM_RE.finditer(content))
    items: list[dict[str, Any]] = []
    for idx, match in enumerate(matches):
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(content)
        title = _clean_outline_title(content[start:end])
        if not title or _is_boilerplate(title):
            continue
        chapter_num, section_num, child_num = match.groups()
        items.append({
            "chapter": int(chapter_num),
            "section": int(section_num),
            "child": int(child_num) if child_num else None,
            "title": title,
        })
    return items


def _find_subsection(chapter: dict[str, Any], section_number: int) -> dict[str, Any] | None:
    subsections = chapter.get("subsections") or []
    if 1 <= section_number <= len(subsections):
        return subsections[section_number - 1]
    return None


def _add_numbered_items(
    chapter: dict[str, Any],
    chapter_number: int,
    text: str,
    warnings: list[str],
    unparsed_items: list[str],
) -> bool:
    added = False
    for item in _split_numbered_items(text):
        if item["chapter"] != chapter_number:
            _add_warning(
                warnings,
                f"Bỏ qua mục {item['chapter']}.{item['section']} vì không khớp Chương {chapter_number}.",
            )
            unparsed_items.append(item["title"])
            continue

        child_number = item.get("child")
        if child_number is None:
            chapter.setdefault("subsections", []).append({
                "title": item["title"],
                "children": [],
            })
            added = True
            continue

        parent = _find_subsection(chapter, item["section"])
        if parent is None:
            _add_warning(
                warnings,
                f"Bỏ qua tiểu mục {item['chapter']}.{item['section']}.{child_number} vì chưa có mục cha.",
            )
            unparsed_items.append(item["title"])
            continue
        parent.setdefault("children", []).append({"title": item["title"]})
        added = True
    return added


def _split_dash_items(line: str) -> list[str]:
    raw = str(line or "").strip()
    if not raw.startswith("-"):
        return []
    matches = re.findall(r"(?:^|\s)-\s+(.+?)(?=\s+-\s+|$)", raw)
    return [_clean_outline_title(match) for match in matches if _clean_outline_title(match)]


def _split_dash_item_records(line: str) -> list[dict[str, Any]]:
    raw = str(line or "").strip()
    if not raw.startswith("-"):
        return []
    records: list[dict[str, Any]] = []
    for match in re.findall(r"(?:^|\s)-\s+(.+?)(?=\s+-\s+|$)", raw):
        title = _clean_outline_title(match)
        if title:
            records.append({
                "title": title,
                "expects_children": match.rstrip().endswith(":"),
            })
    return records


def _add_bullets(
    chapter: dict[str, Any],
    lines: list[str],
    warnings: list[str],
    unparsed_items: list[str],
) -> bool:
    added = False
    last_subsection: dict[str, Any] | None = None
    collecting_children = False
    for raw_line in lines:
        line = str(raw_line or "").strip()
        if not line:
            continue
        if _is_boilerplate(line):
            collecting_children = False
            continue
        if line.startswith("-"):
            collecting_children = False
            for item in _split_dash_item_records(line):
                title = item["title"]
                if _is_boilerplate(title):
                    continue
                last_subsection = {"title": title, "children": []}
                chapter.setdefault("subsections", []).append(last_subsection)
                collecting_children = bool(item["expects_children"])
                added = True
            continue
        child_match = re.match(r"^[•*◦]\s+(.+)$", line)
        if child_match:
            title = _clean_outline_title(child_match.group(1))
            if not title or _is_boilerplate(title):
                continue
            if last_subsection is None:
                _add_warning(warnings, "Có bullet con nhưng chưa nhận diện được mục cha.")
                unparsed_items.append(title)
                continue
            last_subsection.setdefault("children", []).append({"title": title})
            added = True
            continue
        if collecting_children and last_subsection is not None:
            title = _clean_outline_title(line)
            if title and not _is_boilerplate(title):
                last_subsection.setdefault("children", []).append({"title": title})
                added = True
    return added


def _structure_depth(curriculum: dict[str, Any]) -> str:
    for chapter in curriculum.get("chapters") or []:
        for subsection in chapter.get("subsections") or []:
            if subsection.get("children"):
                return "level2"
    return "level1"


def _finalize_parsed_curriculum(
    *,
    topic: str,
    chapters: list[dict[str, Any]],
    target_pages: int | None,
    warnings: list[str],
) -> dict[str, Any]:
    valid_chapters: list[dict[str, Any]] = []
    for idx, chapter in enumerate(chapters, start=1):
        chapter_title = _clean_outline_title(chapter.get("title"))
        subsections = [
            {
                **subsection,
                "title": _clean_outline_title(subsection.get("title")),
                "children": [
                    {"title": _clean_outline_title(child.get("title"))}
                    for child in subsection.get("children") or []
                    if _clean_outline_title(child.get("title"))
                    and not _is_boilerplate(child.get("title"))
                ],
            }
            for subsection in chapter.get("subsections") or []
            if _clean_outline_title(subsection.get("title"))
            and not _is_boilerplate(subsection.get("title"))
        ]
        if not chapter_title or not subsections:
            _add_warning(
                warnings,
                f"Chương {idx} không đủ tiêu đề hoặc mục nên không được tự điền.",
            )
            continue
        cleaned_chapter: dict[str, Any] = {
            "title": chapter_title,
            "subsections": subsections,
        }
        if chapter.get("target_pages"):
            cleaned_chapter["target_pages"] = chapter["target_pages"]
        valid_chapters.append(cleaned_chapter)

    if not valid_chapters:
        raise StructureParseError(
            "Không nhận diện được cấu trúc giáo trình. Vui lòng gửi file có bảng TT / Nội dung / Số trang, "
            "ví dụ: Chương 1 | Tên chương | 10 và các dòng 1.1 | Tên mục."
        )

    if target_pages is None:
        chapter_pages = [
            chapter.get("target_pages")
            for chapter in valid_chapters
            if isinstance(chapter.get("target_pages"), int)
        ]
        if len(chapter_pages) == len(valid_chapters) and chapter_pages:
            target_pages = sum(chapter_pages)
            _add_warning(warnings, "Không thấy dòng Tổng số trang; hệ thống tạm dùng tổng số trang theo chương.")

    curriculum: dict[str, Any] = {
        "topic": topic,
        "chapters": valid_chapters,
    }
    if target_pages is not None:
        curriculum["target_pages"] = target_pages
    curriculum["structure_depth"] = _structure_depth(curriculum)
    return curriculum


def _parse_table_rows(
    rows: list[dict[str, Any]],
    *,
    topic: str,
    warnings: list[str],
    unparsed_items: list[str],
) -> dict[str, Any]:
    chapters: list[dict[str, Any]] = []
    current_chapter: dict[str, Any] | None = None
    current_chapter_number: int | None = None
    target_pages: int | None = None

    for row in rows:
        cells = row.get("cells") or []
        cell_lines = row.get("cell_lines") or []
        if not any(str(cell or "").strip() for cell in cells):
            continue

        normalized_row = " ".join(_normalize_ascii(cell) for cell in cells)
        if "tong so trang" in normalized_row:
            target_pages = _parse_page_count(cells[-1] if cells else "")
            continue
        if "noi dung" in normalized_row and "so trang" in normalized_row:
            continue

        first_cell = str(cells[0] if cells else "").strip()
        content_cell = str(cells[1] if len(cells) > 1 else first_cell).strip()
        page_cell = cells[2] if len(cells) > 2 else ""
        chapter_match = _CHAPTER_RE.match(first_cell) or _CHAPTER_RE.match(content_cell)

        if chapter_match:
            current_chapter_number = int(chapter_match.group(1))
            chapter_title = _chapter_title_from_content(content_cell, current_chapter_number)
            if not chapter_title:
                chapter_title = _clean_outline_title(_strip_chapter_prefix(first_cell, current_chapter_number))
            current_chapter = {
                "title": chapter_title,
                "subsections": [],
            }
            chapter_pages = _parse_page_count(page_cell)
            if chapter_pages is not None:
                current_chapter["target_pages"] = chapter_pages
            chapters.append(current_chapter)
            _add_numbered_items(
                current_chapter,
                current_chapter_number,
                content_cell,
                warnings,
                unparsed_items,
            )
            continue

        if current_chapter is None or current_chapter_number is None:
            continue

        section_numbered_text = ""
        first_cell_number = _NUMBERED_ITEM_RE.match(f"{first_cell} ")
        if first_cell_number:
            section_numbered_text = f"{first_cell} {content_cell}"
        else:
            section_numbered_text = content_cell

        added_numbered = _add_numbered_items(
            current_chapter,
            current_chapter_number,
            section_numbered_text,
            warnings,
            unparsed_items,
        )
        content_lines = cell_lines[1] if len(cell_lines) > 1 else content_cell.splitlines()
        added_bullets = _add_bullets(current_chapter, content_lines, warnings, unparsed_items)

        if not added_numbered and not added_bullets and content_cell and not _is_boilerplate(content_cell):
            if not current_chapter.get("subsections"):
                for line in content_lines:
                    title = _clean_outline_title(line)
                    if title and not _is_boilerplate(title):
                        current_chapter.setdefault("subsections", []).append({
                            "title": title,
                            "children": [],
                        })
                        added_bullets = True
                if added_bullets:
                    continue
            unparsed_items.append(content_cell)

    return _finalize_parsed_curriculum(
        topic=topic,
        chapters=chapters,
        target_pages=target_pages,
        warnings=warnings,
    )


def _extract_topic(lines: list[str]) -> str:
    for idx, line in enumerate(lines):
        match = _TOPIC_RE.search(line)
        if match:
            title = _clean_outline_title(match.group(1))
            if title:
                return title
            if idx + 1 < len(lines):
                return _clean_outline_title(lines[idx + 1])
    return ""


def _docx_tables(content: bytes) -> tuple[list[list[dict[str, Any]]], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise StructureParseError("Backend chưa cài python-docx để đọc file Word.") from exc

    document = Document(BytesIO(content))
    all_lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[list[dict[str, Any]]] = []
    for table in document.tables:
        rows: list[dict[str, Any]] = []
        for row in table.rows:
            cell_lines = []
            cells = []
            for cell in row.cells:
                lines = [
                    line.strip()
                    for paragraph in cell.paragraphs
                    for line in paragraph.text.splitlines()
                    if line.strip()
                ]
                cell_lines.append(lines)
                cells.append(" ".join(lines))
                all_lines.extend(lines)
            rows.append({"cells": cells, "cell_lines": cell_lines})
        tables.append(rows)
    return tables, all_lines


def _pdf_lines(content: bytes) -> list[str]:
    try:
        import fitz  # type: ignore

        document = fitz.open(stream=content, filetype="pdf")
        return [
            line.strip()
            for page in document
            for line in page.get_text("text").splitlines()
            if line.strip()
        ]
    except ImportError:
        pass
    except Exception as exc:
        raise StructureParseError(f"Không đọc được file PDF: {exc}") from exc

    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        return [
            line.strip()
            for page in reader.pages
            for line in (page.extract_text() or "").splitlines()
            if line.strip()
        ]
    except Exception as exc:
        raise StructureParseError(f"Không đọc được file PDF: {exc}") from exc


def _text_rows(lines: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for line in lines:
        parts = [part.strip() for part in re.split(r"\s*\|\s*|\t+", line) if part.strip()]
        if len(parts) >= 2:
            rows.append({"cells": parts[:3], "cell_lines": [[part] for part in parts[:3]]})
        else:
            rows.append({"cells": ["", line, ""], "cell_lines": [[], [line], []]})
    return rows


def _is_page_only_line(line: str) -> bool:
    return bool(re.fullmatch(r"~?\s*\d{1,4}\s*", str(line or "").strip()))


def _is_standalone_section_number(line: str) -> bool:
    return bool(re.fullmatch(r"\d{1,2}\.\d{1,2}", str(line or "").strip()))


def _is_pdf_marker_line(line: str) -> bool:
    normalized = _normalize_ascii(line)
    return (
        not normalized
        or normalized in {"tt", "noi dung", "so trang"}
        or normalized.startswith("mau gt")
        or normalized.startswith("ii. cau truc")
        or normalized.startswith("iii. tai lieu")
        or normalized.startswith("bo giao duc")
        or normalized.startswith("truong dai hoc")
        or normalized.startswith("cong hoa xa hoi")
        or normalized.startswith("doc lap")
        or normalized.startswith("can tho, ngay")
        or normalized.startswith("hieu truong")
        or normalized.startswith("phong ql nckh")
        or normalized.startswith("truong khoa")
        or normalized.startswith("chu bien")
        or normalized.startswith("(ky")
    )


def _extract_trailing_page(text: str) -> tuple[str, str]:
    raw = str(text or "").strip()
    match = re.search(r"(.*?)\s+(~?\d{1,4})\s*$", raw)
    if not match:
        return raw, ""
    title, page = match.groups()
    if not title.strip():
        return raw, ""
    return title.strip(), page.strip()


def _merge_fragmented_pdf_lines(lines: list[str]) -> list[str]:
    merged: list[str] = []
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        normalized = _normalize_ascii(line)
        if idx + 1 < len(lines) and normalized in {"bai thuc", "bai thuc hanh"}:
            next_line = lines[idx + 1]
            next_normalized = _normalize_ascii(next_line)
            if re.match(r"^(?:hanh\s*)?\d+\b", next_normalized):
                merged.append(f"{line} {next_line}")
                idx += 2
                continue
        merged.append(line)
        idx += 1
    return merged


def _pdf_rows_from_lines(lines: list[str]) -> list[dict[str, Any]]:
    """
    Convert text extracted from Word-exported PDFs into table-like rows.

    Word-to-PDF conversion often emits table cells as sequential lines rather
    than preserving real table structure. This keeps the downstream row parser
    identical for DOCX and PDF.
    """
    rows: list[dict[str, Any]] = []
    cleaned = _merge_fragmented_pdf_lines([
        str(line or "").strip()
        for line in lines
        if str(line or "").strip()
    ])
    structure_start = next(
        (
            line_idx + 1
            for line_idx, line in enumerate(cleaned)
            if "cau truc" in _normalize_ascii(line)
            and "noi dung" in _normalize_ascii(line)
        ),
        0,
    )
    if structure_start:
        cleaned = cleaned[structure_start:]
    idx = 0
    while idx < len(cleaned):
        line = cleaned[idx]
        normalized = _normalize_ascii(line)

        if _is_pdf_marker_line(line) or _is_page_only_line(line):
            idx += 1
            continue

        if "tong so trang" in normalized:
            page = ""
            if idx + 1 < len(cleaned) and _is_page_only_line(cleaned[idx + 1]):
                page = cleaned[idx + 1]
                idx += 1
            else:
                _, page = _extract_trailing_page(line)
            rows.append({
                "cells": ["Tổng số trang", "Tổng số trang", page],
                "cell_lines": [["Tổng số trang"], ["Tổng số trang"], [page] if page else []],
            })
            idx += 1
            continue

        chapter_match = _CHAPTER_RE.match(line)
        if chapter_match:
            first_cell = line
            content = _strip_chapter_prefix(line, int(chapter_match.group(1)))
            page = ""

            if content:
                content, page = _extract_trailing_page(content)
            if not content and idx + 1 < len(cleaned):
                next_line = cleaned[idx + 1]
                next_normalized = _normalize_ascii(next_line)
                if (
                    not _CHAPTER_RE.match(next_line)
                    and not _is_standalone_section_number(next_line)
                    and "tong so trang" not in next_normalized
                    and not _is_pdf_marker_line(next_line)
                ):
                    content = next_line
                    idx += 1
                    if idx + 1 < len(cleaned) and _is_page_only_line(cleaned[idx + 1]):
                        page = cleaned[idx + 1]
                        idx += 1
            elif idx + 1 < len(cleaned) and _is_page_only_line(cleaned[idx + 1]):
                page = cleaned[idx + 1]
                idx += 1

            rows.append({
                "cells": [first_cell, content, page],
                "cell_lines": [[first_cell], content.splitlines() if content else [], [page] if page else []],
            })
            idx += 1
            continue

        if _is_standalone_section_number(line):
            content_lines: list[str] = []
            idx += 1
            while idx < len(cleaned):
                next_line = cleaned[idx]
                next_normalized = _normalize_ascii(next_line)
                if (
                    _CHAPTER_RE.match(next_line)
                    or _is_standalone_section_number(next_line)
                    or "tong so trang" in next_normalized
                ):
                    break
                if _is_page_only_line(next_line) or _is_pdf_marker_line(next_line):
                    idx += 1
                    continue
                content_lines.append(next_line)
                idx += 1
            content = " ".join(content_lines)
            rows.append({
                "cells": [line, content, ""],
                "cell_lines": [[line], content_lines, []],
            })
            continue

        if line.startswith("-"):
            detail_lines = [line]
            idx += 1
            while idx < len(cleaned):
                next_line = cleaned[idx]
                next_normalized = _normalize_ascii(next_line)
                if (
                    _CHAPTER_RE.match(next_line)
                    or _is_standalone_section_number(next_line)
                    or "tong so trang" in next_normalized
                ):
                    break
                if not _is_pdf_marker_line(next_line):
                    detail_lines.append(next_line)
                idx += 1
            rows.append({
                "cells": ["", "\n".join(detail_lines), ""],
                "cell_lines": [[], detail_lines, []],
            })
            continue

        rows.append({"cells": ["", line, ""], "cell_lines": [[], [line], []]})
        idx += 1

    return rows


def parse_structure_document(filename: str, content: bytes) -> dict[str, Any]:
    """Parse a Word/PDF outline file into the structured-editor curriculum shape."""
    if not content:
        raise StructureParseError("File rỗng. Vui lòng gửi lại file đề cương có nội dung.")
    if len(content) > _MAX_STRUCTURE_FILE_SIZE:
        raise StructureParseError("File vượt quá giới hạn 10MB. Vui lòng gửi file đề cương gọn hơn.")

    clean_filename = str(filename or "").lower()
    warnings: list[str] = []
    unparsed_items: list[str] = []

    if clean_filename.endswith(".docx"):
        tables, lines = _docx_tables(content)
        topic = _extract_topic(lines)
        for rows in tables:
            header_found = any(
                "noi dung" in " ".join(_normalize_ascii(cell) for cell in row.get("cells") or [])
                and "so trang" in " ".join(_normalize_ascii(cell) for cell in row.get("cells") or [])
                for row in rows
            )
            if not header_found:
                continue
            curriculum = _parse_table_rows(
                rows,
                topic=topic,
                warnings=warnings,
                unparsed_items=unparsed_items,
            )
            return {
                "topic": topic,
                "target_pages": curriculum.get("target_pages"),
                "structure_depth": curriculum.get("structure_depth", "level1"),
                "curriculum": curriculum,
                "warnings": warnings,
                "unparsed_items": unparsed_items[:30],
                "source_format": "docx",
            }
        _add_warning(warnings, "Không tìm thấy bảng TT / Nội dung / Số trang; đã thử nhận dạng từ văn bản.")
        rows = _text_rows(lines)
        source_format = "docx"
    elif clean_filename.endswith(".pdf"):
        lines = _pdf_lines(content)
        topic = _extract_topic(lines)
        _add_warning(
            warnings,
            "PDF chỉ hỗ trợ file có text chọn được; file scan ảnh cần gửi lại dạng Word hoặc PDF text.",
        )
        rows = _pdf_rows_from_lines(lines)
        source_format = "pdf"
    else:
        raise StructureParseError("Chỉ hỗ trợ file .docx hoặc .pdf cho chức năng tự điền cấu trúc.")

    try:
        curriculum = _parse_table_rows(
            rows,
            topic=topic,
            warnings=warnings,
            unparsed_items=unparsed_items,
        )
    except StructureParseError:
        if source_format != "pdf":
            raise
        warnings = [
            "PDF chỉ hỗ trợ file có text chọn được; file scan ảnh cần gửi lại dạng Word hoặc PDF text.",
            "Không khôi phục được bảng từ PDF; đã thử nhận dạng từ văn bản phẳng.",
        ]
        unparsed_items = []
        curriculum = _parse_table_rows(
            _text_rows(lines),
            topic=topic,
            warnings=warnings,
            unparsed_items=unparsed_items,
        )
    return {
        "topic": topic,
        "target_pages": curriculum.get("target_pages"),
        "structure_depth": curriculum.get("structure_depth", "level1"),
        "curriculum": curriculum,
        "warnings": warnings,
        "unparsed_items": unparsed_items[:30],
        "source_format": source_format,
    }


def _clean_heading_text(value: Any) -> str:
    return str(value or "").strip()


def clean_subsection_title(value: Any) -> str:
    title = _clean_heading_text(value)
    cleaned = _SUBSECTION_PREFIX_RE.sub("", title).strip()
    return cleaned or title


def parse_structure_markdown(markdown: str, topic: str = "") -> dict[str, Any]:
    """
    Parse a Markdown outline where ``##`` is a chapter, ``###`` is a level-1
    subsection, and optional ``####`` headings are level-2 controlled leaves.

    Non-empty non-heading lines are rejected so the contract remains predictable
    for both the UI serializer and external API callers.
    """
    chapters: list[dict[str, Any]] = []
    current_chapter: dict[str, Any] | None = None
    current_subsection: dict[str, Any] | None = None

    for line_number, raw_line in enumerate(str(markdown or "").splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue

        match = _HEADING_RE.match(line)
        if not match:
            raise StructureParseError(
                f"Line {line_number} must be a ## chapter or ### subsection heading"
            )

        marker, raw_title = match.groups()
        title = _clean_heading_text(raw_title)
        if not title:
            raise StructureParseError(f"Line {line_number} has an empty heading title")

        if marker == "##":
            current_chapter = {"title": title, "subsections": []}
            current_subsection = None
            chapters.append(current_chapter)
            continue

        if current_chapter is None:
            raise StructureParseError("A ### subsection cannot appear before a ## chapter")

        if marker == "###":
            subsection_title = clean_subsection_title(title)
            current_subsection = {"title": subsection_title}
            current_chapter["subsections"].append(current_subsection)
            continue

        if current_subsection is None:
            raise StructureParseError("A #### child subsection cannot appear before a ### subsection")

        child_title = clean_subsection_title(title)
        current_subsection.setdefault("children", []).append({"title": child_title})

    if not chapters:
        raise StructureParseError("Structure must contain at least one chapter")

    for idx, chapter in enumerate(chapters, start=1):
        if not chapter["title"]:
            raise StructureParseError(f"Chapter {idx} title cannot be empty")
        if not chapter["subsections"]:
            raise StructureParseError(f"Chapter {idx} must contain at least one subsection")

    return {
        "topic": _clean_heading_text(topic),
        "chapters": chapters,
    }


def structure_to_markdown(curriculum: dict[str, Any]) -> str:
    """Serialize a curriculum-like dict to the Markdown outline contract."""
    lines: list[str] = []
    chapters = curriculum.get("chapters") if isinstance(curriculum, dict) else []
    for chapter_index, chapter in enumerate(chapters or [], start=1):
        chapter_title = _clean_heading_text(chapter.get("title"))
        if not chapter_title:
            continue
        lines.append(f"## {chapter_title}")
        subsections = chapter.get("subsections") or []
        for subsection_index, subsection in enumerate(subsections, start=1):
            subsection_title = _clean_heading_text(subsection.get("title"))
            if subsection_title:
                lines.append(f"### {chapter_index}.{subsection_index} {subsection_title}")
                for child_index, child in enumerate(subsection.get("children") or [], start=1):
                    child_title = _clean_heading_text(child.get("title"))
                    if child_title:
                        lines.append(
                            f"#### {chapter_index}.{subsection_index}.{child_index} {child_title}"
                        )
    return "\n".join(lines)
