import pytest
from io import BytesIO

from app.services.textbook.structure_parser import (
    StructureParseError,
    parse_structure_document,
    parse_structure_markdown,
    structure_to_markdown,
)


def _docx_bytes(rows, topic="TEST GIÁO TRÌNH"):
    from docx import Document

    document = Document()
    document.add_paragraph(f"- Tên giáo trình: {topic}")
    table = document.add_table(rows=len(rows), cols=3)
    for row_idx, row_values in enumerate(rows):
        for col_idx, value in enumerate(row_values):
            table.cell(row_idx, col_idx).text = value
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _weekly_docx_bytes(rows, topic="Xử lý ảnh"):
    from docx import Document

    document = Document()
    document.add_paragraph(f"Tên học phần: {topic}")
    document.add_paragraph("Nội dung chi tiết học phần")
    table = document.add_table(rows=len(rows), cols=4)
    for row_idx, row_values in enumerate(rows):
        for col_idx, value in enumerate(row_values):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            for line_idx, line in enumerate(str(value).splitlines()):
                paragraph = cell.paragraphs[0] if line_idx == 0 else cell.add_paragraph()
                paragraph.text = line
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_parse_structure_markdown_uses_chapters_and_subsections() -> None:
    curriculum = parse_structure_markdown(
        """
## Chapter One
### 1.1 First section
### (1.2) Second section
## Chapter Two
### Mục 2.1 Third section
""",
        topic="Test topic",
    )

    assert curriculum["topic"] == "Test topic"
    assert [chapter["title"] for chapter in curriculum["chapters"]] == [
        "Chapter One",
        "Chapter Two",
    ]
    assert [sub["title"] for sub in curriculum["chapters"][0]["subsections"]] == [
        "First section",
        "Second section",
    ]
    assert curriculum["chapters"][1]["subsections"][0]["title"] == "Third section"


def test_parse_structure_rejects_subsection_before_chapter() -> None:
    with pytest.raises(StructureParseError):
        parse_structure_markdown("### 1.1 No chapter")


def test_parse_structure_rejects_empty_chapter() -> None:
    with pytest.raises(StructureParseError):
        parse_structure_markdown("## Empty chapter")


def test_structure_to_markdown_round_trip_contract() -> None:
    markdown = structure_to_markdown({
        "chapters": [
            {
                "title": "Chapter",
                "subsections": [{"title": "Section"}],
            }
        ]
    })

    assert markdown == "## Chapter\n### 1.1 Section"


def test_parse_structure_markdown_supports_level2_children() -> None:
    curriculum = parse_structure_markdown(
        """
## Chapter One
### 1.1 Parent section
#### 1.1.1 Child A
#### 1.1.2 Child B
""",
        topic="Nested topic",
    )

    subsection = curriculum["chapters"][0]["subsections"][0]
    assert subsection["title"] == "Parent section"
    assert [child["title"] for child in subsection["children"]] == ["Child A", "Child B"]


def test_structure_to_markdown_serializes_level2_children() -> None:
    markdown = structure_to_markdown({
        "chapters": [
            {
                "title": "Chapter",
                "subsections": [
                    {
                        "title": "Parent",
                        "children": [{"title": "Child"}],
                    }
                ],
            }
        ]
    })

    assert markdown == "## Chapter\n### 1.1 Parent\n#### 1.1.1 Child"


def test_parse_structure_document_reads_standard_docx_table() -> None:
    result = parse_structure_document(
        "outline.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Chương 1", "Tổng quan mạng máy tính", "15"],
            ["1.1", "Khái niệm mạng máy tính", ""],
            ["1.2", "Mô hình OSI", ""],
            ["Tổng số trang", "Tổng số trang", "15"],
        ], topic="MẠNG MÁY TÍNH"),
    )

    assert result["topic"] == "MẠNG MÁY TÍNH"
    assert result["target_pages"] == 15
    chapter = result["curriculum"]["chapters"][0]
    assert chapter["title"] == "Tổng quan mạng máy tính"
    assert chapter["target_pages"] == 15
    assert [sub["title"] for sub in chapter["subsections"]] == [
        "Khái niệm mạng máy tính",
        "Mô hình OSI",
    ]


def test_parse_structure_document_splits_inline_numbered_items() -> None:
    result = parse_structure_document(
        "outline.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            [
                "Chương 2",
                "Chương 2: Active Directory 2.1. Các mô hình mạng 2.1.1. Workgroup 2.1.2. Domain 2.2. Dịch vụ AD Câu hỏi ôn tập",
                "~15",
            ],
            ["Tổng số trang", "Tổng số trang", "15"],
        ]),
    )

    chapter = result["curriculum"]["chapters"][0]
    assert chapter["title"] == "Active Directory"
    assert chapter["target_pages"] == 15
    assert result["structure_depth"] == "level2"
    assert chapter["subsections"][0]["title"] == "Các mô hình mạng"
    assert [child["title"] for child in chapter["subsections"][0]["children"]] == [
        "Workgroup",
        "Domain",
    ]
    assert chapter["subsections"][1]["title"] == "Dịch vụ AD"


def test_parse_structure_document_treats_lesson_and_dash_rows_as_chapter_sections() -> None:
    result = parse_structure_document(
        "practice.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Bài 1", "Cài đặt Windows Server\nTóm tắt chương", "12"],
            ["", "- Cài đặt trên máy ảo\n- Cấu hình IP tĩnh\n- Cập nhật hệ thống\nCâu hỏi trắc nghiệm ôn tập", ""],
            ["Tổng số trang", "Tổng số trang", "12"],
        ]),
    )

    chapter = result["curriculum"]["chapters"][0]
    assert chapter["title"] == "Cài đặt Windows Server"
    assert [sub["title"] for sub in chapter["subsections"]] == [
        "Cài đặt trên máy ảo",
        "Cấu hình IP tĩnh",
        "Cập nhật hệ thống",
    ]


def test_parse_structure_document_reads_weekly_detail_table_as_level1_sections() -> None:
    result = parse_structure_document(
        "course-outline.docx",
        _weekly_docx_bytes([
            ["Tuần", "Nội dung", "Tài liệu", "CĐR của HP"],
            ["1", "Chương 1 – Tổng quan về xử lý ảnh", "[1] [2]", "CO1"],
            [
                "1",
                "Giới thiệu lịch sử hình thành\nMột số thuật ngữ\nCác bước xử lý ảnh",
                "[1] [2]",
                "CO1",
            ],
            [
                "2,3",
                "Chương 2 – Xử lý điểm trên ảnh\n- Giới thiệu\n- Xử lý Histogram\n- Các phép toán số học",
                "[1] [2]",
                "CO2",
            ],
            [
                "10",
                "Báo cáo nhóm\nChủ đề báo cáo cuối kỳ",
                "[1] [2]",
                "CO1",
            ],
        ]),
    )

    assert result["topic"] == "Xử lý ảnh"
    assert result["source_format"] == "docx_weekly_detail"
    chapters = result["curriculum"]["chapters"]
    assert [chapter["title"] for chapter in chapters] == [
        "Tổng quan về xử lý ảnh",
        "Xử lý điểm trên ảnh",
    ]
    assert [sub["title"] for sub in chapters[0]["subsections"]] == [
        "Giới thiệu lịch sử hình thành",
        "Một số thuật ngữ",
        "Các bước xử lý ảnh",
    ]
    assert [sub["title"] for sub in chapters[1]["subsections"]] == [
        "Giới thiệu",
        "Xử lý Histogram",
        "Các phép toán số học",
    ]
    assert "Báo cáo nhóm" in result["unparsed_items"]


def test_parse_structure_document_treats_lines_after_colon_dash_as_children() -> None:
    result = parse_structure_document(
        "practice.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Bài 11", "Giám sát truy cập, in ấn và theo dõi hệ thống\nTóm tắt chương", "12"],
            [
                "",
                "- Thiết lập Auditing\n- Cấu hình Printer Server\n- Theo dõi hệ thống:\nTheo dõi hiệu suất\nGiám sát máy chủ và dịch vụ\nTheo dõi nhật ký hệ thống\nCâu hỏi trắc nghiệm ôn tập",
                "",
            ],
            ["Tổng số trang", "Tổng số trang", "12"],
        ]),
    )

    chapter = result["curriculum"]["chapters"][0]
    assert result["structure_depth"] == "level2"
    assert [sub["title"] for sub in chapter["subsections"]] == [
        "Thiết lập Auditing",
        "Cấu hình Printer Server",
        "Theo dõi hệ thống",
    ]
    assert [child["title"] for child in chapter["subsections"][2]["children"]] == [
        "Theo dõi hiệu suất",
        "Giám sát máy chủ và dịch vụ",
        "Theo dõi nhật ký hệ thống",
    ]


def test_parse_structure_document_treats_lesson_and_numbered_rows_as_chapter_sections() -> None:
    result = parse_structure_document(
        "practice.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Bài 1", "Phân tích gói tin mạng với Wireshark [2] [5]", "12"],
            ["", "1.1 Giới thiệu Wireshark. 1.2 Hướng dẫn bắt gói tin. 1.3 Ôn tập lý thuyết TCP/IP.", ""],
            ["Tổng số trang", "Tổng số trang", "12"],
        ]),
    )

    chapter = result["curriculum"]["chapters"][0]
    assert chapter["title"] == "Phân tích gói tin mạng với Wireshark"
    assert [sub["title"] for sub in chapter["subsections"]] == [
        "Giới thiệu Wireshark",
        "Hướng dẫn bắt gói tin",
        "Ôn tập lý thuyết TCP/IP",
    ]


def test_parse_structure_document_treats_practice_lesson_label_as_chapter() -> None:
    result = parse_structure_document(
        "practice.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Bài thực\nhành 1", "Xử lý dữ liệu", "15"],
            ["1.1", "Các phương pháp tiền xử lý và xử lý dữ liệu", ""],
            ["1.2", "Bài tập thực hành", ""],
            ["Tổng số trang", "Tổng số trang", "15"],
        ], topic="THỰC HÀNH KHOA HỌC DỮ LIỆU"),
    )

    assert result["topic"] == "THỰC HÀNH KHOA HỌC DỮ LIỆU"
    chapter = result["curriculum"]["chapters"][0]
    assert chapter["title"] == "Xử lý dữ liệu"
    assert chapter["target_pages"] == 15
    assert [sub["title"] for sub in chapter["subsections"]] == [
        "Các phương pháp tiền xử lý và xử lý dữ liệu",
        "Bài tập thực hành",
    ]


def test_parse_structure_document_rejects_unsupported_or_empty_files() -> None:
    with pytest.raises(StructureParseError):
        parse_structure_document("outline.txt", "TT | Nội dung | Số trang".encode("utf-8"))
    with pytest.raises(StructureParseError):
        parse_structure_document("outline.docx", b"")


def test_parse_structure_document_returns_partial_result_with_warnings() -> None:
    result = parse_structure_document(
        "outline.docx",
        _docx_bytes([
            ["TT", "Nội dung", "Số trang"],
            ["Chương 1", "Chương thiếu mục", "10"],
            ["Chương 2", "Chương dùng được", "10"],
            ["2.1", "Mục dùng được", ""],
        ]),
    )

    assert [chapter["title"] for chapter in result["curriculum"]["chapters"]] == ["Chương dùng được"]
    assert result["warnings"]


def test_parse_structure_document_groups_word_exported_pdf_lines(monkeypatch) -> None:
    import app.services.textbook.structure_parser as parser

    monkeypatch.setattr(parser, "_pdf_lines", lambda _content: [
        "- Tên giáo trình:",
        "THỰC HÀNH QUẢN TRỊ MẠNG",
        "TT",
        "Nội dung",
        "Số trang",
        "Bài 11",
        "Giám sát truy cập, in ấn và theo dõi hệ thống",
        "12",
        "Tóm tắt chương",
        "- Thiết lập Auditing",
        "- Theo dõi hệ thống:",
        "Theo dõi hiệu suất",
        "Giám sát máy chủ và dịch vụ",
        "Câu hỏi trắc nghiệm ôn tập",
        "Bài 12",
        "Hướng dẫn báo cáo hết môn",
        "10",
        "1.1",
        "Trình bày yêu cầu báo cáo",
        "Tổng số trang",
        "22",
    ])

    result = parse_structure_document("outline.pdf", b"pdf bytes")

    assert result["source_format"] == "pdf"
    assert result["topic"] == "THỰC HÀNH QUẢN TRỊ MẠNG"
    assert result["target_pages"] == 22
    assert result["structure_depth"] == "level2"
    chapters = result["curriculum"]["chapters"]
    assert [chapter["title"] for chapter in chapters] == [
        "Giám sát truy cập, in ấn và theo dõi hệ thống",
        "Hướng dẫn báo cáo hết môn",
    ]
    assert chapters[0]["subsections"][1]["title"] == "Theo dõi hệ thống"
    assert [child["title"] for child in chapters[0]["subsections"][1]["children"]] == [
        "Theo dõi hiệu suất",
        "Giám sát máy chủ và dịch vụ",
    ]
    assert chapters[1]["subsections"][0]["title"] == "Trình bày yêu cầu báo cáo"


def test_parse_structure_document_groups_fragmented_practice_lesson_pdf_lines(monkeypatch) -> None:
    import app.services.textbook.structure_parser as parser

    monkeypatch.setattr(parser, "_pdf_lines", lambda _content: [
        "1",
        "Mẫu GT 1.1",
        "- Tên giáo trình: THỰC HÀNH KHOA HỌC DỮ LIỆU",
        "II. Cấu trúc - Nội dung giáo trình",
        "TT",
        "Nội dung",
        "Số trang",
        "Bài thực",
        "hành 1",
        "Xử lý dữ liệu",
        "15",
        "1.1",
        "Các phương pháp tiền xử lý và xử lý dữ liệu: Làm sạch, loại bỏ",
        "nhiễu, outliers, kết hợp dữ liệu, chuấn hóa dữ liệu...",
        "BỘ GIÁO DỤC VÀ ĐÀO TẠO",
        "TRƯỜNG ĐẠI HỌC NAM CẦN THƠ",
        "2",
        "1.2",
        "Bài tập thực hành",
        "Bài thực",
        "hành 2",
        "Dự đoán dữ liệu",
        "15",
        "2.1",
        "Cài đặt các thuật toán phân lớp: Decision Tree, SVM, k-NN,",
        "Naïve Bayes...",
        "Tổng số trang",
        "30",
    ])

    result = parse_structure_document("practice.pdf", b"pdf bytes")

    assert result["source_format"] == "pdf"
    assert result["topic"] == "THỰC HÀNH KHOA HỌC DỮ LIỆU"
    assert result["target_pages"] == 30
    chapters = result["curriculum"]["chapters"]
    assert [chapter["title"] for chapter in chapters] == ["Xử lý dữ liệu", "Dự đoán dữ liệu"]
    assert chapters[0]["target_pages"] == 15
    assert chapters[0]["subsections"][0]["title"] == (
        "Các phương pháp tiền xử lý và xử lý dữ liệu: Làm sạch, loại bỏ "
        "nhiễu, outliers, kết hợp dữ liệu, chuấn hóa dữ liệu"
    )
    assert chapters[0]["subsections"][1]["title"] == "Bài tập thực hành"
    assert chapters[1]["subsections"][0]["title"] == (
        "Cài đặt các thuật toán phân lớp: Decision Tree, SVM, k-NN, Naïve Bayes"
    )


def test_structured_planner_preserves_user_structure(monkeypatch) -> None:
    from app.services.textbook.planner import HybridPlanner

    planner = HybridPlanner.__new__(HybridPlanner)

    def fake_enrich(*args, **kwargs):
        return [
            {
                "description": "Description A",
                "search_query": "query a",
                "section_type": "medium",
            },
            {
                "description": "Description B",
                "search_query": "query b",
                "section_type": "applied",
            },
        ]

    monkeypatch.setattr(planner, "_enrich_structured_chapter", fake_enrich)
    curriculum = planner.enrich_user_structure(
        "Topic",
        "Requirements",
        {
            "chapters": [
                {
                    "title": "Chapter A",
                    "subsections": [
                        {"title": "Section A"},
                        {"title": "Section B"},
                    ],
                }
            ]
        },
        language="en",
    )

    assert curriculum is not None
    assert curriculum.chapters[0].title == "Chapter A"
    assert [sub.title for sub in curriculum.chapters[0].subsections] == [
        "Section A",
        "Section B",
    ]
    assert curriculum.chapters[0].subsections[1].section_type == "applied"


def test_structured_planner_preserves_nested_user_structure(monkeypatch) -> None:
    from app.services.textbook.planner import HybridPlanner

    planner = HybridPlanner.__new__(HybridPlanner)

    def fake_enrich(*args, **kwargs):
        return [
            {
                "description": "Child A description",
                "search_query": "query child a",
                "section_type": "medium",
            },
            {
                "description": "Child B description",
                "search_query": "query child b",
                "section_type": "applied",
            },
        ]

    monkeypatch.setattr(planner, "_enrich_structured_chapter", fake_enrich)
    curriculum = planner.enrich_user_structure(
        "Topic",
        "Requirements",
        {
            "chapters": [
                {
                    "title": "Chapter A",
                    "subsections": [
                        {
                            "title": "Parent A",
                            "target_pages": 6,
                            "children": [
                                {"title": "Child A", "target_pages": 3},
                                {"title": "Child B", "target_pages": 3},
                            ],
                        },
                    ],
                }
            ]
        },
        language="en",
        structure_depth="level2",
    )

    assert curriculum is not None
    assert curriculum.structure_depth == "level2"
    parent = curriculum.chapters[0].subsections[0]
    assert parent.title == "Parent A"
    assert parent.target_pages == 6
    assert [child.title for child in parent.children] == ["Child A", "Child B"]
    assert parent.children[1].section_type == "applied"
    assert parent.children[0].target_pages == 3
